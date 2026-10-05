# -*- coding: utf-8 -*-
"""judge verification: rubric checks on package/out vs oracle/out, case1 x4."""
import json, sys, io, zipfile, re, subprocess, shutil, os
from pathlib import Path
from lxml import etree
import docx
from docx.oxml.ns import qn

ROOT = Path(r"D:\workspace\zcode研究\skillfactory\v3\assets\office-templates")
PKG, ORA, INP = ROOT/"package/out", ROOT/"oracle/out", ROOT/"oracle/inputs"
W = Path(r"D:\workspace\zcode研究\_tmp_judge_ab1_rep1")
TS = [u"周报", u"请示函", u"会议通知", u"工作总结"]
FIELDS = {
 u"周报": ([u"部门",u"填报人",u"周期",u"本周工作内容",u"下周工作计划",u"问题与需协调事项",u"报送日期"],
           [1,1,1,1,1,0,0], [u"list",u"list",u"list"] , [3,4,5]),
 u"请示函": ([u"请示事由",u"主送机关",u"请示缘由",u"请示事项",u"请示单位",u"联系人",u"联系电话",u"成文日期"],
           [1,1,1,1,1,0,0,0], [], []),
 u"会议通知": ([u"会议名称",u"召开单位",u"主送对象",u"会议时间",u"会议地点",u"参会人员",u"会议议题",u"会议要求",u"联系人",u"联系电话",u"发文日期"],
           [1,1,1,1,1,1,0,0,0,0,0], [], [7]),
 u"工作总结": ([u"总结主体",u"总结时段",u"工作回顾",u"主要成绩",u"存在问题",u"下一步工作打算",u"成文日期"],
           [1,1,1,1,0,1,0], [], [2,3,4,5]),
}
SPEC_TOP = {u"template",u"title",u"filled_fields",u"missing_fields",u"fields",u"summary",u"outputs"}
SPEC_SUM = {u"total",u"filled",u"missing",u"missing_required",u"missing_required_names",u"unknown_keys"}
fails, notes = [], []
def chk(cond, msg):
    (notes if cond else fails).append(("PASS " if cond else "FAIL ")+msg)

def para_props(p):
    """return dict of alignment, ind attrs, spacing attrs for a python-docx paragraph."""
    pPr = p._p.pPr
    out = {"jc": None, "ind": {}, "spacing": {}}
    if pPr is None: return out
    jc = pPr.find(qn("w:jc"))
    if jc is not None: out["jc"] = jc.get(qn("w:val"))
    ind = pPr.find(qn("w:ind"))
    if ind is not None:
        out["ind"] = {k.split('}')[-1]: v for k, v in ind.attrib.items()}
    sp = pPr.find(qn("w:spacing"))
    if sp is not None:
        out["spacing"] = {k.split('}')[-1]: v for k, v in sp.attrib.items()}
    return out

def run_props(p):
    r = []
    for run in p.runs:
        rPr = run._r.rPr
        d = {"eastAsia": None, "ascii": None, "hAnsi": None, "sz": None, "bold": None}
        if rPr is not None:
            rf = rPr.find(qn("w:rFonts"))
            if rf is not None:
                d["eastAsia"] = rf.get(qn("w:eastAsia")); d["ascii"] = rf.get(qn("w:ascii")); d["hAnsi"] = rf.get(qn("w:hAnsi"))
            sz = rPr.find(qn("w:sz"))
            if sz is not None: d["sz"] = int(sz.get(qn("w:val")))/2.0
            b = rPr.find(qn("w:b"))
            if b is not None: d["bold"] = b.get(qn("w:val"), "1")
        r.append(d)
    return r

def sect(doc):
    s = doc.sections[0]
    EMU_CM = 360000.0
    return {"w": s.page_width.cm, "h": s.page_height.cm,
            "top": s.top_margin.cm, "bot": s.bottom_margin.cm,
            "left": s.left_margin.cm, "right": s.right_margin.cm}

def strip_outputs(obj):
    o = json.loads(json.dumps(obj)); o.pop("outputs", None); return o

def canon(obj):
    return json.dumps(strip_outputs(obj), ensure_ascii=False, indent=2, sort_keys=False)

report = {}
for t in TS:
    rep = {}
    inp = json.loads((INP/t/"case1.json").read_text(encoding="utf-8-sig"))
    fj_p = json.loads((PKG/t/"case1"/"fields.json").read_text(encoding="utf-8"))
    fj_o = json.loads((ORA/t/"case1"/"fields.json").read_text(encoding="utf-8"))
    # A. fields.json schema
    rep["top_keys_ok"] = set(fj_p.keys()) == SPEC_TOP
    rep["sum_keys_ok"] = set(fj_p["summary"].keys()) == SPEC_SUM
    names, req, kinds, listidx = FIELDS[t]
    det = fj_p["fields"]
    rep["field_names_ok"] = [d["name"] for d in det] == names
    rep["item_keys_ok"] = all(set(d.keys()) >= {"name","required","kind","status","value"} and set(d.keys()) <= {"name","required","kind","status","value","rendered_as"} for d in det)
    rep["req_kind_ok"] = all(d["required"] == bool(req[i]) and d["kind"] == (u"list" if i in listidx else u"text") for i, d in enumerate(det))
    filled = [d["name"] for d in det if d["status"]=="filled"]; missing = [d["name"] for d in det if d["status"]=="missing"]
    rep["s3_lists_ok"] = filled == fj_p["filled_fields"] and missing == fj_p["missing_fields"]
    s = fj_p["summary"]
    rep["s3_counts_ok"] = (s["total"]==len(det)==len(names) and s["filled"]==len(filled) and s["missing"]==len(missing)
        and s["filled"]+s["missing"]==s["total"]
        and s["missing_required_names"]==[d["name"] for d in det if d["status"]=="missing" and d["required"]]
        and s["missing_required"]==len(s["missing_required_names"]))
    rep["filled_all"] = (s["filled"], s["total"])
    rep["title_ok"] = fj_p["title"] == (docx.Document(str(PKG/t/"case1"/u"文书.docx")).paragraphs[0].text if True else "")
    # title formula vs input
    if t == u"周报": tf = inp[u"部门"]+u"工作周报"
    elif t == u"请示函": tf = u"关于"+inp[u"请示事由"]+u"的请示"
    elif t == u"会议通知": tf = u"关于召开"+inp[u"会议名称"]+u"的通知"
    else: tf = inp[u"总结主体"]+inp[u"总结时段"]+u"工作总结"
    rep["title_formula_ok"] = (fj_p["title"] == tf, fj_p["title"], tf)
    # fields byte-equal to oracle (modulo outputs)
    rep["fields_eq_oracle_bytes"] = canon(fj_p) == canon(fj_o)
    # B. docx layout
    d = docx.Document(str(PKG/t/"case1"/u"文书.docx"))
    paras = [p for p in d.paragraphs if p.text.strip()]
    g = sect(d)
    rep["page"] = {k: round(v,3) for k,v in g.items()}
    rep["v1_ok"] = (abs(g["w"]-21.0)<=0.05 and abs(g["h"]-29.7)<=0.05 and abs(g["top"]-3.7)<=0.05
        and abs(g["bot"]-3.5)<=0.05 and abs(g["left"]-2.8)<=0.05 and abs(g["right"]-2.6)<=0.05)
    ttl = paras[0]; tp = para_props(ttl); tr = run_props(ttl)
    rep["v2"] = {"jc": tp["jc"], "eastAsia": tr[0]["eastAsia"] if tr else None, "sz": tr[0]["sz"] if tr else None,
                 "text_eq_title": ttl.text == fj_p["title"]}
    # salutation
    if t in (u"请示函", u"会议通知"):
        sal = paras[1]; sp_ = para_props(sal)
        flc = sp_["ind"].get("firstLineChars")
        key = u"主送机关" if t==u"请示函" else u"主送对象"
        rep["v3"] = {"text": sal.text, "flush": flc in (None,"0"), "ends_colon": sal.text.strip().endswith(u"："),
                     "eq_input": sal.text.strip() == inp[key].strip()+u"："}
    else:
        rep["v3"] = "n/a"
    # body: any para with firstLineChars=200 and proper runs/spacing
    body_ok, body_detail = False, []
    for p in paras[1:]:
        pr = para_props(p)
        if pr["ind"].get("firstLineChars") == "200":
            r = run_props(p)
            ea = {x["eastAsia"] for x in r}; asc = {x["ascii"] for x in r}; han = {x["hAnsi"] for x in r}; szs = {x["sz"] for x in r}
            sp2 = pr["spacing"]
            okr = (ea=={u"仿宋"} and asc<= {u"Times New Roman", None} and han <= {u"Times New Roman", None} and szs=={16.0}
                   and sp2.get("line")=="560" and sp2.get("lineRule")=="exact" and sp2.get("before") in (None,"0") and sp2.get("after") in (None,"0"))
            body_detail.append((p.text[:18], sorted(ea), sorted(szs), sp2.get("line"), sp2.get("lineRule"), okr))
            if okr: body_ok = True
    rep["v4_ok"] = body_ok; rep["v4_detail_n"] = len(body_detail)
    # signature
    sig = paras[-1]; sp3 = para_props(sig)
    right = sp3["ind"].get("right")
    rep["v5"] = {"text": sig.text[:20], "jc": sp3["jc"], "right": right, "right_pt": int(right)/20.0 if right else None}
    # content-diff pins vs oracle
    do = docx.Document(str(ORA/t/"case1"/u"文书.docx"))
    txt_p = "\n".join(p.text for p in d.paragraphs); txt_o = "\n".join(p.text for p in do.paragraphs)
    pins = {u"周报": [u"部门：", u"填报人："], u"请示函": [u"现就有关事项请示如下", u"妥否，请批示"],
            u"会议通知": [u"经研究，决定召开", u"特此通知", u"一、会议要求", u"五、"]}.get(t, [])
    rep["pins"] = {pin: {"A": pin in txt_p, "B": pin in txt_o} for pin in pins}
    rep["para_count"] = {"A": len([p for p in d.paragraphs if p.text.strip()]), "B": len([p for p in do.paragraphs if p.text.strip()])}
    # body firstLine vs firstLineChars (arm-b L2 claim) — inspect ind attrs of A body paras
    ind_attrs = set()
    for p in paras[1:]:
        pr = para_props(p)
        ind_attrs.add(tuple(sorted(pr["ind"].items())))
    rep["A_ind_attr_variants"] = [dict(x) for x in list(ind_attrs)[:6]]
    report[t] = rep

# styles.xml docDefaults + title ascii font (arm-b L1/L4 claims)
def docdefaults(path):
    with zipfile.ZipFile(path) as z:
        s = z.read("word/styles.xml").decode("utf-8")
    return u"w:docDefaults" in s, s[:0]
for t in TS:
    a = PKG/t/"case1"/u"文书.docx"; b = ORA/t/"case1"/u"文书.docx"
    rep = report[t]
    rep["L1_docDefaults"] = {"A": docdefaults(a)[0], "B": docdefaults(b)[0]}
    # title run ascii font
    da = docx.Document(str(a)); ta = [p for p in da.paragraphs if p.text.strip()][0]
    db = docx.Document(str(b)); tb = [p for p in db.paragraphs if p.text.strip()][0]
    ra, rb = run_props(ta)[0], run_props(tb)[0]
    rep["L4_title_ascii"] = {"A": (ra["ascii"], ra["hAnsi"]), "B": (rb["ascii"], rb["hAnsi"])}

for t in TS:
    print("="*8, t)
    for k, v in report[t].items():
        print(f"  {k}: {v}")
print("="*20)
print("FAILS:", len(fails))
for f in fails: print(" ", f)
