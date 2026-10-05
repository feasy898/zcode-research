# -*- coding: utf-8 -*-
"""Build GMP 审计缺陷整改会 meeting minutes: 纪要.docx + summary.json + OUT.md (in temp dir)."""
import json
import sys
from pathlib import Path

from docx import Document

TMP = Path(__file__).resolve().parent

# ---- source transcript lines (from the ask, verbatim) ----
LINES = {
    "1102": "[11:02] 生产部2车间主任: 关键偏差的调查报告还没写完，这个跟进我来负责，下周一之前提交质量部。",
    "1104": "[11:04] 注册事务部小齐: 变更备案的递交截止到监管平台关闭日，具体时间待通知。",
    "1106": "[11:06] 质量部总监: 有个风险，审计缺陷如果带着过，下次飞检大概率升级，存在停产整改的不确定因素。",
    "1108": "[11:08] 生产部2车间主任: 等等，我刚才说的提交时间不对，是下下周三之前，不是下周一。",
    "1110": "[11:10] 质量部总监: 结论：缺陷整改实行清单化管理，一条一销号。",
}

def sent(key: str) -> str:
    """Original sentence body (strip '[hh:mm] speaker: ' prefix)."""
    return LINES[key].split(": ", 1)[1]

# verbatim-quote sanity: each sentence must be an exact substring of its line
for k, line in LINES.items():
    assert sent(k) in line, k

# ---- minutes content (single source of truth) ----
TITLE = "GMP审计缺陷整改会 会议纪要"
H1, H2, H3 = "一、决议事项", "二、待办事项", "三、风险与关注"

RESOLUTIONS = [
    f"1. 【11:10 · 质量部总监】{sent('1110')}",
]
TODOS = [
    {
        "事项": f"【11:02 · 生产部2车间主任】{sent('1102')}",
        "负责人": "生产部2车间主任",  # 原句「这个跟进我来负责」→ 说话人本人
        "期限": "下下周三之前",        # 原句「下周一之前」已被【11:08】明确更正
    },
    {
        "事项": f"【11:04 · 注册事务部小齐】{sent('1104')}",
        "负责人": "待定",  # 原句无明确负责人指派
        "期限": "待定",    # 原句明示「具体时间待通知」，无可确定日期
    },
]
RISKS = [
    f"1. 【11:06 · 质量部总监】{sent('1106')}",
]
NOTE = (
    "注：待办第1项的期限以发言人在【11:08 · 生产部2车间主任】"
    f"「{sent('1108')}」中的更正为准（原句所述「下周一之前」作废）；"
    "待办第2项因原句明示「具体时间待通知」且未指派负责人，按确定性规则填「待定」。"
)

COUNTS = {"决议事项": len(RESOLUTIONS), "待办事项": len(TODOS), "风险与关注": len(RISKS)}
COUNTS["合计"] = COUNTS["决议事项"] + COUNTS["待办事项"] + COUNTS["风险与关注"]

# ---- 纪要.docx ----
doc = Document()
doc.add_heading(TITLE, level=0)

doc.add_heading(H1, level=1)
for item in RESOLUTIONS:
    doc.add_paragraph(item)

doc.add_heading(H2, level=1)
table = doc.add_table(rows=1 + len(TODOS), cols=3)
table.style = "Table Grid"
for j, h in enumerate(["事项", "负责人", "期限"]):
    table.rows[0].cells[j].text = h
for i, todo in enumerate(TODOS, start=1):
    for j, key in enumerate(["事项", "负责人", "期限"]):
        table.rows[i].cells[j].text = todo[key]
doc.add_paragraph(NOTE)

doc.add_heading(H3, level=1)
for item in RISKS:
    doc.add_paragraph(item)

doc.save(TMP / "纪要.docx")

# ---- summary.json ----
(TMP / "summary.json").write_text(
    json.dumps(COUNTS, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)

# ---- OUT.md ----
md = []
md.append(f"# {TITLE}\n")
md.append(f"## {H1}\n")
for item in RESOLUTIONS:
    md.append(item + "\n")
md.append(f"## {H2}\n")
md.append("| 事项 | 负责人 | 期限 |")
md.append("| --- | --- | --- |")
for t in TODOS:
    md.append(f"| {t['事项']} | {t['负责人']} | {t['期限']} |")
md.append("")
md.append(f"> {NOTE}\n")
md.append(f"## {H3}\n")
for item in RISKS:
    md.append(item + "\n")
(TMP / "OUT.md").write_text("\n".join(md), encoding="utf-8")

print("BUILT:", sorted(p.name for p in TMP.iterdir()))
print("COUNTS:", COUNTS)
