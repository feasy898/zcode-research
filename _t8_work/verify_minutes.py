# -*- coding: utf-8 -*-
"""Verify built artifacts before delivering: docx structure, verbatim quotes, summary counts."""
import json
import re
from pathlib import Path

from docx import Document

TMP = Path(__file__).resolve().parent

TRANSCRIPT = [
    "[11:00] 质量部总监: 上午好，人到齐了开始，今天过GMP审计缺陷的整改。",
    "[11:02] 生产部2车间主任: 关键偏差的调查报告还没写完，这个跟进我来负责，下周一之前提交质量部。",
    "[11:04] 注册事务部小齐: 变更备案的递交截止到监管平台关闭日，具体时间待通知。",
    "[11:06] 质量部总监: 有个风险，审计缺陷如果带着过，下次飞检大概率升级，存在停产整改的不确定因素。",
    "[11:08] 生产部2车间主任: 等等，我刚才说的提交时间不对，是下下周三之前，不是下周一。",
    "[11:10] 质量部总监: 结论：缺陷整改实行清单化管理，一条一销号。",
    "[11:12] 质量部总监: 散会，各位辛苦。",
]

doc = Document(TMP / "纪要.docx")
paras = [p.text for p in doc.paragraphs]
fail = []

def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        fail.append(msg)

# 1. three section headings, exact text, in order
idx = [paras.index(h) if h in paras else -1 for h in ("一、决议事项", "二、待办事项", "三、风险与关注")]
check(all(i >= 0 for i in idx), f"three exact headings present: {idx}")
check(idx == sorted(idx) and -1 not in idx, "headings in order 1/2/3")

# 2. resolutions & risks are numbered items quoting original sentences verbatim with prefix
def section_texts(a, b):
    return [t for t in paras[idx[a] + 1 : idx[b]]] if b is not None else paras[idx[a] + 1 :]

res_items = [t for t in section_texts(0, 1) if t.strip()]
risk_items = [t for t in section_texts(2, None) if t.strip()]
check(len(res_items) == 1, f"resolution item count == 1 (got {len(res_items)})")
check(len(risk_items) == 1, f"risk item count == 1 (got {len(risk_items)})")
check(all(re.match(r"^\d+\.\s", t) for t in res_items + risk_items), "决议/风险 items are numbered")

def verbatim_prefixed(item_text):
    m = re.match(r"^(?:\d+\.\s*)?【(\d{2}:\d{2}) · ([^】]+)】(.+)$", item_text)
    if not m:
        return False
    t, sp, body = m.groups()
    for line in TRANSCRIPT:
        lm = re.match(r"^\[(\d{2}:\d{2})\] ([^:]+): (.+)$", line)
        if lm and lm.group(1) == t and lm.group(2) == sp and lm.group(3) == body:
            return True
    return False

check(all(verbatim_prefixed(t) for t in res_items + risk_items), "决议/风险 quotes verbatim + 【时间 · 说话人】prefix")

# 3. table: 3 columns, header 事项/负责人/期限, 2 data rows
check(len(doc.tables) == 1, f"exactly one table (got {len(doc.tables)})")
tb = doc.tables[0]
check(len(tb.columns) == 3, f"table has 3 columns (got {len(tb.columns)})")
check([c.text for c in tb.rows[0].cells] == ["事项", "负责人", "期限"], "header cells == 事项/负责人/期限")
rows = [[c.text for c in r.cells] for r in tb.rows[1:]]
check(len(rows) == 2, f"2 todo rows (got {len(rows)})")
for r in rows:
    check(verbatim_prefixed(r[0]), f"todo 事项 verbatim+prefix: {r[0][:20]}...")
check(rows[0][1] == "生产部2车间主任" and rows[0][2] == "下下周三之前", f"row1 owner/deadline: {rows[0][1]}/{rows[0][2]}")
check(rows[1][1] == "待定" and rows[1][2] == "待定", f"row2 owner/deadline: {rows[1][1]}/{rows[1][2]}")

# 4. summary.json: exactly four keys, ints, matches docx, 合计 = sum
s = json.loads((TMP / "summary.json").read_text(encoding="utf-8"))
check(list(s.keys()) == ["决议事项", "待办事项", "风险与关注", "合计"], f"summary keys: {list(s.keys())}")
check(all(isinstance(v, int) for v in s.values()), "summary values are ints")
check((s["决议事项"], s["待办事项"], s["风险与关注"]) == (len(res_items), len(rows), len(risk_items)),
      f"summary counts match docx: {s['决议事项']}/{s['待办事项']}/{s['风险与关注']} vs docx {len(res_items)}/{len(rows)}/{len(risk_items)}")
check(s["合计"] == s["决议事项"] + s["待办事项"] + s["风险与关注"], f"合计 == sum: {s['合计']}")

# 5. OUT.md contains all body content
md = (TMP / "OUT.md").read_text(encoding="utf-8")
for h in ("一、决议事项", "二、待办事项", "三、风险与关注"):
    check(f"## {h}" in md, f"OUT.md has heading {h}")
for t in res_items + risk_items:
    check(t in md, f"OUT.md contains item: {t[:20]}...")
for r in rows:
    check(all(cell in md for cell in r), f"OUT.md contains todo row cells: {r[1]}/{r[2]}")

print("\n" + ("ALL CHECKS PASSED" if not fail else f"{len(fail)} CHECK(S) FAILED"))
raise SystemExit(1 if fail else 0)
