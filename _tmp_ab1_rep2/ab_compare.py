# -*- coding: utf-8 -*-
"""A/B comparator: package/out vs oracle/out (case1, four templates).
Reads paired artifacts, prints ONLY derived findings (diffs / layout profiles / digests).
"""
import hashlib, json, sys, io
from docx import Document
from docx.oxml.ns import qn

BASE = r"D:\workspace\zcode研究\skillfactory\v3\assets\office-templates"
BAK  = r"D:\workspace\zcode研究\_tmp_ab1_rep2\backup2"
TMPLS = ["周报", "请示函", "会议通知", "工作总结"]

def sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()

def load_json(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)

def json_diff(a, b, path=""):
    out = []
    if isinstance(a, dict) and isinstance(b, dict):
        for k in sorted(set(a) | set(b)):
            if k not in a:
                out.append(f"{path}/{k}: ONLY_IN_PACKAGE = {b[k]!r}")
            elif k not in b:
                out.append(f"{path}/{k}: ONLY_IN_ORACLE = {a[k]!r}")
            else:
                out += json_diff(a[k], b[k], f"{path}/{k}")
    elif isinstance(a, list) and isinstance(b, list):
        if len(a) != len(b):
            out.append(f"{path}: LIST_LEN {len(a)} vs {len(b)}")
        for i, (x, y) in enumerate(zip(a, b)):
            out += json_diff(x, y, f"{path}[{i}]")
    else:
        if a != b:
            out.append(f"{path}: VALUE oracle={a!r} vs package={b!r}")
    return out

def run_info(r):
    rpr = r._element.rPr
    east = None
    if rpr is not None and rpr.rFonts is not None:
        east = rpr.rFonts.get(qn("w:eastAsia"))
    sz = r.font.size.pt if r.font.size is not None else None
    return {"t": r.text, "bold": r.bold, "italic": r.italic, "underline": r.underline,
            "ascii_font": r.font.name, "east_font": east, "size_pt": sz}

def para_info(p):
    pf = p.paragraph_format
    return {"style": p.style.name if p.style is not None else None,
            "align": str(pf.alignment),
            "space_before": pf.space_before.pt if pf.space_before is not None else None,
            "space_after": pf.space_after.pt if pf.space_after is not None else None,
            "line_spacing": str(pf.line_spacing),
            "first_line_indent": pf.first_line_indent.pt if pf.first_line_indent is not None else None,
            "text": p.text, "runs": [run_info(r) for r in p.runs]}

def docx_profile(path):
    d = Document(path)
    prof = {"sections": [], "paras": [para_info(p) for p in d.paragraphs], "tables": []}
    for s in d.sections:
        prof["sections"].append({
            "page_w_cm": round(s.page_width.cm, 3) if s.page_width else None,
            "page_h_cm": round(s.page_height.cm, 3) if s.page_height else None,
            "orientation": str(s.orientation),
            "margins_cm": [round(x.cm, 3) if x is not None else None
                           for x in (s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)]})
    for t in d.tables:
        prof["tables"].append({
            "style": t.style.name if t.style is not None else None,
            "rows": len(t.rows), "cols": len(t.columns),
            "cells": [[c.text for c in row.cells] for row in t.rows]})
    # also default style of Normal
    try:
        n = d.styles["Normal"]
        prof["normal_style"] = {"font": n.font.name, "size_pt": n.font.size.pt if n.font.size else None}
    except Exception as e:
        prof["normal_style"] = f"ERR {e}"
    return prof

def profile_diff(a, b, path=""):
    """Structural diff of two profiles; text contents compared too."""
    out = []
    if isinstance(a, dict) and isinstance(b, dict):
        for k in sorted(set(a) | set(b)):
            if k not in a:
                out.append(f"{path}.{k}: ONLY_IN_PACKAGE")
            elif k not in b:
                out.append(f"{path}.{k}: ONLY_IN_ORACLE")
            else:
                out += profile_diff(a[k], b[k], f"{path}.{k}")
    elif isinstance(a, list) and isinstance(b, list):
        if len(a) != len(b):
            out.append(f"{path}: LEN oracle={len(a)} package={len(b)}")
        for i, (x, y) in enumerate(zip(a, b)):
            out += profile_diff(x, y, f"{path}[{i}]")
    else:
        if a != b:
            out.append(f"{path}: oracle={a!r} vs package={b!r}")
    return out

def brief_layout(prof, tag):
    print(f"  [{tag}] sections={prof['sections']}")
    print(f"  [{tag}] normal_style={prof['normal_style']}")
    print(f"  [{tag}] paragraphs={len(prof['paras'])} tables={len(prof['tables'])}")
    for i, p in enumerate(prof["paras"]):
        runs = "; ".join(
            f"<{r['t'][:18]!r} b={r['bold']} font={r['ascii_font']}/{r['east_font']} {r['size_pt']}pt>"
            for r in p["runs"]) or "(no runs)"
        print(f"  [{tag}] P{i:02d} style={p['style']} align={p['align']} "
              f"indent={p['first_line_indent']} text={p['text'][:26]!r} runs: {runs}")
    for i, t in enumerate(prof["tables"]):
        print(f"  [{tag}] T{i} style={t['style']} {t['rows']}x{t['cols']} "
              f"row0={t['cells'][0] if t['cells'] else None}")

for T in TMPLS:
    print("=" * 100)
    print(f"TEMPLATE {T}  (case1)")
    pk_docx = rf"{BASE}\package\out\{T}\case1\文书.docx"
    or_docx = rf"{BASE}\oracle\out\{T}\case1\文书.docx"
    pk_json = rf"{BASE}\package\out\{T}\case1\fields.json"
    or_json = rf"{BASE}\oracle\out\{T}\case1\fields.json"

    # --- fields.json ledger ---
    oj, pj = load_json(or_json), load_json(pk_json)
    print(f"-- fields.json: oracle sha={sha(or_json)[:12]} package sha={sha(pk_json)[:12]}")
    dj = json_diff(oj, pj)
    print(f"-- fields.json deep-diff: {'IDENTICAL' if not dj else str(len(dj)) + ' diff(s)'}")
    for line in dj:
        print("     " + line)
    print(f"-- oracle ledger top-level keys: {list(oj.keys()) if isinstance(oj, dict) else type(oj)}")

    # --- docx ---
    same_bytes = sha(or_docx) == sha(pk_docx)
    print(f"-- 文书.docx: oracle sha={sha(or_docx)[:12]} package sha={sha(pk_docx)[:12]} byte_identical={same_bytes}")
    po, pp = docx_profile(or_docx), docx_profile(pk_docx)
    dp = profile_diff(po, pp)
    print(f"-- docx layout profile diff: {'LAYOUT IDENTICAL' if not dp else str(len(dp)) + ' diff(s)'}")
    for line in dp:
        print("     " + line)
    brief_layout(pp, "PKG")
    if dp:
        brief_layout(po, "ORC")

    # --- determinism: fresh outputs vs pre-run backup ---
    for tag, fresh, old in (("oracle", or_docx, rf"{BAK}\oracle_out\{T}\case1\文书.docx"),
                            ("oracle", or_json, rf"{BAK}\oracle_out\{T}\case1\fields.json"),
                            ("package", pk_docx, rf"{BAK}\package_out\{T}\case1\文书.docx"),
                            ("package", pk_json, rf"{BAK}\package_out\{T}\case1\fields.json")):
        import os
        exists = os.path.exists(old)
        det = exists and sha(fresh) == sha(old)
        print(f"-- determinism {tag} {os.path.basename(old)}: backup_exists={exists} fresh==backup:{det}")
print("=" * 100)
print("DONE")
