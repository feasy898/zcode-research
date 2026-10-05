#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""oracle.py —— 四类中文办公文书参照生成器（oracle）。

用 python-docx 按中文公文/事务文书规范版式生成 docx：
  * 标题居中（黑体二号）
  * 称谓顶格（无首行缩进）
  * 正文首行缩进两字符（w:firstLineChars=200，三号仿宋，固定行距 28 磅）
  * 落款右对齐（右空两字）
  * A4 页面、版心按 GB/T 9704-2012（上 3.7 / 下 3.5 / 左 2.8 / 右 2.6 cm）

用法：
  python oracle.py --template <周报|请示函|会议通知|工作总结> --data <json> --outdir <dir>

产物（写入 outdir）：
  文书.docx    生成的文书
  fields.json  实际填充字段与留空字段清单

字段语义：
  * 必填字段缺失/为空      -> 文中以 "____" 占位（列表字段渲染占位条目），fields.json 记入 missing
  * 选填字段缺失/为空      -> 行内字段留空（空字符串）；独立条目/小节省略渲染，fields.json 记入 missing
  * 数据中未被模板消费的键 -> fields.json 的 summary.unknown_keys
  * 产物成功写出即退出码 0（必填缺失属数据边界，如实记录而非报错）；
    模板名非法 / 数据文件缺失 / JSON 非法 / 数据非对象 -> 退出码 2
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

try:
    from docx.oxml.parser import OxmlElement
except ImportError:  # 兼容老版本 python-docx
    from docx.oxml import OxmlElement

# ---- 版式常量 ----
TITLE_FONT = "黑体"
BODY_FONT = "仿宋"
HEAD_FONT = "黑体"
ASCII_FONT = "Times New Roman"
BODY_SIZE = Pt(16)      # 三号
TITLE_SIZE = Pt(22)     # 二号
LINE_SPACING = Pt(28)   # 固定值 28 磅
PLACEHOLDER = "____"

CENTER = WD_ALIGN_PARAGRAPH.CENTER
RIGHT = WD_ALIGN_PARAGRAPH.RIGHT

CN_NUM = "一二三四五六七八九十"


def cn_num(i: int) -> str:
    return CN_NUM[i - 1] if 1 <= i <= len(CN_NUM) else str(i)


class FieldTracker:
    """逐字段记录实际取值：填充（filled）或留空（missing）。"""

    def __init__(self, data: dict):
        self.data = data
        self.filled: dict = {}
        self.missing: list = []

    def _record_missing(self, name: str, required: bool, kind: str, rendered_as: str) -> None:
        if all(m["field"] != name for m in self.missing):
            self.missing.append(
                {"field": name, "required": required, "kind": kind, "rendered_as": rendered_as}
            )

    def text(self, name: str, required: bool) -> str:
        raw = self.data.get(name)
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            self._record_missing(name, required, "text",
                                 "以____占位" if required else "留空（渲染为空字符串）")
            return PLACEHOLDER if required else ""
        value = str(raw).strip()
        self.filled[name] = value
        return value

    def items(self, name: str, required: bool) -> list:
        raw = self.data.get(name)
        if raw is None:
            values = []
        elif isinstance(raw, str):
            values = [raw.strip()] if raw.strip() else []
        elif isinstance(raw, list):
            values = [str(x).strip() for x in raw if str(x).strip()]
        else:
            values = [str(raw).strip()]
        if not values:
            self._record_missing(name, required, "list",
                                 "以____占位条目" if required else "省略该条目/小节")
            return [PLACEHOLDER] if required else []
        self.filled[name] = values
        return values


# ---- 版式辅助 ----

def new_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)   # A4
    sec.top_margin, sec.bottom_margin = Cm(3.7), Cm(3.5)   # GB/T 9704 版心
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.6)
    try:  # Normal 样式兜底；每个 run 仍会显式设字体
        normal = doc.styles["Normal"]
        normal.font.name = ASCII_FONT
        normal.font.size = BODY_SIZE
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
    except Exception:
        pass
    return doc


def _style_run(run, east: str, size, bold: bool) -> None:
    run.font.name = ASCII_FONT
    run.font.size = size
    run.font.bold = bold
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ASCII_FONT)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT)
    rfonts.set(qn("w:eastAsia"), east)


def _set_first_line_chars(p, chars: int) -> None:
    """首行缩进 N 字符：写 firstLineChars（字符数，Word 优先）+ firstLine（twips 兜底）。"""
    p.paragraph_format.first_line_indent = Pt(16 * chars)
    ppr = p._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars * 100))


def para(doc, text: str, *, align=None, east: str = BODY_FONT, size=None, bold: bool = False,
         first_line_chars: int = 0, right_indent_chars: int = 0) -> None:
    size = size or BODY_SIZE
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    if right_indent_chars:
        pf.right_indent = Pt(16 * right_indent_chars)
    if first_line_chars:
        _set_first_line_chars(p, first_line_chars)
    _style_run(p.add_run(text), east=east, size=size, bold=bold)


def _section(doc, heading: str, items: list, num: int | None = None) -> None:
    para(doc, f"{cn_num(num)}、{heading}" if num else heading,
         first_line_chars=2, east=HEAD_FONT)
    for i, item in enumerate(items, 1):
        para(doc, f"{i}．{item}", first_line_chars=2)


# ---- 四类模板 ----

def build_zhoubao(doc: Document, t: FieldTracker) -> str:
    dept = t.text("部门", required=True)
    person = t.text("填报人", required=True)
    period = t.text("周期", required=True)
    date = t.text("报送日期", required=False)

    para(doc, f"{dept}工作周报", align=CENTER, east=TITLE_FONT, size=TITLE_SIZE)  # 标题居中
    para(doc, f"填报人：{person}　　部门：{dept}　　周期：{period}", align=CENTER)
    sections = [
        ("本周工作内容", t.items("本周工作内容", required=True)),
        ("下周工作计划", t.items("下周工作计划", required=True)),
    ]
    problems = t.items("问题与需协调事项", required=False)
    if problems:
        sections.append(("问题与需协调事项", problems))
    for i, (head, items) in enumerate(sections, 1):
        _section(doc, head, items, num=i)
    para(doc, f"报送日期：{date}", align=RIGHT, right_indent_chars=2)  # 落款右对齐
    return f"{dept}工作周报"


def build_qingshihan(doc: Document, t: FieldTracker) -> str:
    reason = t.text("请示事由", required=True)
    superior = t.text("主送机关", required=True)
    basis = t.text("请示缘由", required=True)
    matter = t.text("请示事项", required=True)
    unit = t.text("请示单位", required=True)
    contact = t.text("联系人", required=False)
    phone = t.text("联系电话", required=False)
    date = t.text("成文日期", required=False)

    para(doc, f"关于{reason}的请示", align=CENTER, east=TITLE_FONT, size=TITLE_SIZE)
    para(doc, f"{superior}：")                       # 称谓顶格
    para(doc, basis, first_line_chars=2)
    para(doc, "现就有关事项请示如下：", first_line_chars=2)
    para(doc, matter, first_line_chars=2)
    para(doc, "妥否，请批示。", first_line_chars=2)
    if contact or phone:
        para(doc, f"联系人：{contact}　　联系电话：{phone}", first_line_chars=2)
    para(doc, unit, align=RIGHT, right_indent_chars=2)  # 落款右对齐
    para(doc, date, align=RIGHT, right_indent_chars=2)
    return f"关于{reason}的请示"


def build_huiyitongzhi(doc: Document, t: FieldTracker) -> str:
    name = t.text("会议名称", required=True)
    host = t.text("召开单位", required=True)
    addressee = t.text("主送对象", required=True)
    when = t.text("会议时间", required=True)
    where = t.text("会议地点", required=True)
    who = t.text("参会人员", required=True)
    topic = t.text("会议议题", required=False)
    requirements = t.items("会议要求", required=False)
    contact = t.text("联系人", required=False)
    phone = t.text("联系电话", required=False)
    date = t.text("发文日期", required=False)

    para(doc, f"关于召开{name}的通知", align=CENTER, east=TITLE_FONT, size=TITLE_SIZE)
    para(doc, f"{addressee}：")                      # 称谓顶格
    para(doc, f"经研究，决定召开{name}。现将有关事项通知如下：", first_line_chars=2)
    entries = [("会议时间", when), ("会议地点", where), ("参会人员", who)]
    if topic:
        entries.append(("会议议题", topic))
    if requirements:
        entries.append(("会议要求", "；".join(requirements) + "。"))
    for i, (head, content) in enumerate(entries, 1):
        para(doc, f"{cn_num(i)}、{head}：{content}", first_line_chars=2)
    para(doc, "特此通知。", first_line_chars=2)
    if contact or phone:
        para(doc, f"联系人：{contact}　　联系电话：{phone}", first_line_chars=2)
    para(doc, host, align=RIGHT, right_indent_chars=2)
    para(doc, date, align=RIGHT, right_indent_chars=2)
    return f"关于召开{name}的通知"


def build_gongzuozongjie(doc: Document, t: FieldTracker) -> str:
    subject = t.text("总结主体", required=True)
    period = t.text("总结时段", required=True)
    review = t.items("工作回顾", required=True)
    achievements = t.items("主要成绩", required=True)
    problems = t.items("存在问题", required=False)
    plans = t.items("下一步工作打算", required=True)
    date = t.text("成文日期", required=False)

    para(doc, f"{subject}{period}工作总结", align=CENTER, east=TITLE_FONT, size=TITLE_SIZE)
    sections = [("工作回顾", review), ("主要成绩", achievements)]
    if problems:
        sections.append(("存在问题", problems))
    sections.append(("下一步工作打算", plans))
    for i, (head, items) in enumerate(sections, 1):
        _section(doc, head, items, num=i)
    para(doc, subject, align=RIGHT, right_indent_chars=2)  # 落款右对齐
    para(doc, date, align=RIGHT, right_indent_chars=2)
    return f"{subject}{period}工作总结"


TEMPLATES = {
    "周报": {
        "build": build_zhoubao,
        "fields": [
            {"name": "部门", "required": True, "kind": "text"},
            {"name": "填报人", "required": True, "kind": "text"},
            {"name": "周期", "required": True, "kind": "text"},
            {"name": "本周工作内容", "required": True, "kind": "list"},
            {"name": "下周工作计划", "required": True, "kind": "list"},
            {"name": "问题与需协调事项", "required": False, "kind": "list"},
            {"name": "报送日期", "required": False, "kind": "text"},
        ],
    },
    "请示函": {
        "build": build_qingshihan,
        "fields": [
            {"name": "请示事由", "required": True, "kind": "text"},
            {"name": "主送机关", "required": True, "kind": "text"},
            {"name": "请示缘由", "required": True, "kind": "text"},
            {"name": "请示事项", "required": True, "kind": "text"},
            {"name": "请示单位", "required": True, "kind": "text"},
            {"name": "联系人", "required": False, "kind": "text"},
            {"name": "联系电话", "required": False, "kind": "text"},
            {"name": "成文日期", "required": False, "kind": "text"},
        ],
    },
    "会议通知": {
        "build": build_huiyitongzhi,
        "fields": [
            {"name": "会议名称", "required": True, "kind": "text"},
            {"name": "召开单位", "required": True, "kind": "text"},
            {"name": "主送对象", "required": True, "kind": "text"},
            {"name": "会议时间", "required": True, "kind": "text"},
            {"name": "会议地点", "required": True, "kind": "text"},
            {"name": "参会人员", "required": True, "kind": "text"},
            {"name": "会议议题", "required": False, "kind": "text"},
            {"name": "会议要求", "required": False, "kind": "list"},
            {"name": "联系人", "required": False, "kind": "text"},
            {"name": "联系电话", "required": False, "kind": "text"},
            {"name": "发文日期", "required": False, "kind": "text"},
        ],
    },
    "工作总结": {
        "build": build_gongzuozongjie,
        "fields": [
            {"name": "总结主体", "required": True, "kind": "text"},
            {"name": "总结时段", "required": True, "kind": "text"},
            {"name": "工作回顾", "required": True, "kind": "list"},
            {"name": "主要成绩", "required": True, "kind": "list"},
            {"name": "存在问题", "required": False, "kind": "list"},
            {"name": "下一步工作打算", "required": True, "kind": "list"},
            {"name": "成文日期", "required": False, "kind": "text"},
        ],
    },
}


# ---- fields.json ----

def build_fields_json(template: str, spec_fields: list, tracker: FieldTracker, data: dict,
                      out_docx: Path, out_json: Path, title: str) -> dict:
    spec_names = {f["name"] for f in spec_fields}
    detail = []
    for f in spec_fields:
        name = f["name"]
        if name in tracker.filled:
            detail.append({"name": name, "required": f["required"], "kind": f["kind"],
                           "status": "filled", "value": tracker.filled[name]})
        else:
            m = next((x for x in tracker.missing if x["field"] == name), None)
            detail.append({"name": name, "required": f["required"], "kind": f["kind"],
                           "status": "missing", "value": None,
                           "rendered_as": m["rendered_as"] if m else "未在文书中使用"})
    missing_names = [d["name"] for d in detail if d["status"] == "missing"]
    missing_required = [d["name"] for d in detail if d["status"] == "missing" and d["required"]]
    unknown = sorted(k for k in data.keys() if k not in spec_names)
    return {
        "template": template,
        "title": title,
        "filled_fields": [d["name"] for d in detail if d["status"] == "filled"],
        "missing_fields": missing_names,
        "fields": detail,
        "summary": {
            "total": len(detail),
            "filled": len(detail) - len(missing_names),
            "missing": len(missing_names),
            "missing_required": len(missing_required),
            "missing_required_names": missing_required,
            "unknown_keys": unknown,
        },
        "outputs": {"docx": str(out_docx), "fields_json": str(out_json)},
    }


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="四类中文办公文书参照生成器（oracle）")
    ap.add_argument("--template", required=True, choices=list(TEMPLATES), help="文书类型")
    ap.add_argument("--data", required=True, help="样例数据 JSON 文件路径")
    ap.add_argument("--outdir", required=True, help="产物输出目录")
    args = ap.parse_args(argv)

    data_path = Path(args.data)
    if not data_path.is_file():
        print(f"[oracle] 数据文件不存在: {data_path}", file=sys.stderr)
        return 2
    try:
        data = json.loads(data_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError) as e:
        print(f"[oracle] JSON 解析失败: {e}", file=sys.stderr)
        return 2
    if not isinstance(data, dict):
        print("[oracle] 数据必须是 JSON 对象（键值对）", file=sys.stderr)
        return 2

    tpl = TEMPLATES[args.template]
    tracker = FieldTracker(data)
    doc = new_document()
    title = tpl["build"](doc, tracker)

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    docx_path = outdir / "文书.docx"
    json_path = outdir / "fields.json"
    doc.save(str(docx_path))
    payload = build_fields_json(args.template, tpl["fields"], tracker, data,
                                docx_path, json_path, title)
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
                         encoding="utf-8")

    s = payload["summary"]
    print(f"OK template={args.template} docx={docx_path} fields={json_path} "
          f"filled={s['filled']} missing={s['missing']} missing_required={s['missing_required']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
