#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""gen_doc.py — 四类中文事务文书生成引擎（周报 / 请示函 / 会议通知 / 工作总结）

给定一类文书与一份数据 JSON，按 GB/T 9704 风格版式生成 `文书.docx`，并输出字段
填充台账 `fields.json`（哪些字段填了、哪些缺了、缺的怎么渲染的）。缺数据不编造：
必填缺失以 `____` 占位并如实记账，选填缺失留空或整节省略。

行为语义依据 spec.md（V1–V8 版式 / F1–F6 字段 / S1–S4 台账 / C1–C3 CLI / D1 确定性），
对外接口冻结于 contract.md §2/§3。

用法：
  python gen_doc.py --template <周报|请示函|会议通知|工作总结> --data <数据.json> --outdir <输出目录>

退出码：0 成功（含必填缺失等数据边界）；2 非法输入（模板名非法 / 数据文件不存在 /
JSON 解析失败 / JSON 顶层非对象），此时不产生产物、不建产物目录。
确定性：正文无当前时间、无随机源；同输入连跑两次产物语义逐字节一致（D1）。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

try:  # Windows 控制台缺省 GBK，强制 UTF-8 输出
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001 — 非 Windows / 已重配置时忽略
    pass

# ---------------------------------------------------------------- 常量（冻结口径）

DOCX_NAME = "文书.docx"
FIELDS_NAME = "fields.json"
TEMPLATES = ("周报", "请示函", "会议通知", "工作总结")
PLACEHOLDER = "____"                      # spec V7：必填缺失占位符（四个下划线）
CONTACT_SEP = "　　"                       # 行内字段间隔（两个全角空格）

TITLE_FONT = "黑体"                        # V2 标题 eastAsia 字体
TITLE_SIZE = Pt(22)                        # 二号
BODY_FONT = "仿宋"                         # V4 正文 eastAsia 字体
ASCII_FONT = "Times New Roman"             # V4 正文 ascii/hAnsi 字体
BODY_SIZE = Pt(16)                         # 三号
LINE_SPACING = Pt(28)                      # V4 固定行距（w:line=560, exact）
SIGN_RIGHT_INDENT = Pt(32)                 # V5 落款右缩进两字（≈32pt）

CN_NUMS = ("一", "二", "三", "四", "五", "六", "七", "八", "九", "十")

# spec §3.3 字段表：template -> [(name, required, kind), ...]（顺序冻结）
FIELDS_SPEC = {
    "周报": [("部门", True, "text"), ("填报人", True, "text"), ("周期", True, "text"),
             ("本周工作内容", True, "list"), ("下周工作计划", True, "list"),
             ("问题与需协调事项", False, "list"), ("报送日期", False, "text")],
    "请示函": [("请示事由", True, "text"), ("主送机关", True, "text"), ("请示缘由", True, "text"),
               ("请示事项", True, "text"), ("请示单位", True, "text"),
               ("联系人", False, "text"), ("联系电话", False, "text"), ("成文日期", False, "text")],
    "会议通知": [("会议名称", True, "text"), ("召开单位", True, "text"), ("主送对象", True, "text"),
                 ("会议时间", True, "text"), ("会议地点", True, "text"), ("参会人员", True, "text"),
                 ("会议议题", False, "text"), ("会议要求", False, "list"),
                 ("联系人", False, "text"), ("联系电话", False, "text"), ("发文日期", False, "text")],
    "工作总结": [("总结主体", True, "text"), ("总结时段", True, "text"), ("工作回顾", True, "list"),
                 ("主要成绩", True, "list"), ("存在问题", False, "list"),
                 ("下一步工作打算", True, "list"), ("成文日期", False, "text")],
}
FIELD_NAMES = {t: {n for n, _r, _k in spec} for t, spec in FIELDS_SPEC.items()}

# 缺失渲染语义（spec F1/F2/F4，写入 fields.json 的 rendered_as）
RENDER_REQ_TEXT = "以____占位"
RENDER_OPT_TEXT = "留空（渲染为空字符串）"
RENDER_REQ_LIST = "以____占位条目"
RENDER_OPT_LIST = "省略该条目/小节"


# ---------------------------------------------------------------- 字段规整（F3/F5）

def norm_text(raw) -> tuple[bool, str]:
    """文本字段规整：返回 (是否非缺失, strip 后的值)。
    None → 缺失；字符串 → strip；其他类型 → str().strip()（宽容，F3 精神）。"""
    if raw is None:
        return False, ""
    s = raw.strip() if isinstance(raw, str) else str(raw).strip()
    return bool(s), s


def norm_list(raw) -> tuple[bool, list[str]]:
    """列表字段规整（F3）：字符串 → 单项；数组 → 逐项 str().strip() 后去空；
    其他类型 → 单项。空数组 / 全空白数组 / 空白串 → 缺失。返回 (是否非缺失, 条目)。"""
    if raw is None:
        return False, []
    if isinstance(raw, str):
        s = raw.strip()
        return bool(s), [s] if s else []
    if isinstance(raw, (list, tuple)):
        items = [str(x).strip() for x in raw]
        items = [x for x in items if x]
        return bool(items), items
    s = str(raw).strip()
    return bool(s), [s] if s else []


def pick(values: dict, name: str, required: bool) -> str:
    """取渲染用文本：必填缺失 → ____（F1/V7）；选填缺失 → 空字符串（F2）。"""
    v = values.get(name)
    if v is None:
        return PLACEHOLDER if required else ""
    return v


# ---------------------------------------------------------------- 台账构建（S1–S4）

def build_fields(template: str, data: dict) -> tuple[list[dict], dict]:
    """按字段表逐项记账。返回 (明细列表, 已填值字典 name->value)。"""
    details, values = [], {}
    for name, req, kind in FIELDS_SPEC[template]:
        if kind == "text":
            ok, val = norm_text(data.get(name))
        else:
            ok, val = norm_list(data.get(name))
        if ok:
            details.append({"name": name, "required": req, "kind": kind,
                            "status": "filled", "value": val})
            values[name] = val
        else:
            rendered = (RENDER_REQ_TEXT if req else RENDER_OPT_TEXT) if kind == "text" \
                else (RENDER_REQ_LIST if req else RENDER_OPT_LIST)
            details.append({"name": name, "required": req, "kind": kind,
                            "status": "missing", "value": None, "rendered_as": rendered})
    return details, values


def build_title(template: str, values: dict) -> str:
    """标题文案公式（spec V2），缺失必填字段以 ____ 代入。"""
    if template == "周报":
        return "%s工作周报" % pick(values, "部门", True)
    if template == "请示函":
        return "关于%s的请示" % pick(values, "请示事由", True)
    if template == "会议通知":
        return "关于召开%s的通知" % pick(values, "会议名称", True)
    return "%s%s工作总结" % (pick(values, "总结主体", True), pick(values, "总结时段", True))


def build_meta(template: str, data: dict, title: str,
               docx_path: Path, fields_path: Path) -> tuple[dict, dict]:
    """构建 fields.json 全量对象（S1–S4）并返回 (meta, values)。"""
    details, values = build_fields(template, data)
    filled = [d["name"] for d in details if d["status"] == "filled"]
    missing = [d["name"] for d in details if d["status"] == "missing"]
    missing_required = [d["name"] for d in details
                        if d["status"] == "missing" and d["required"]]
    unknown = sorted(k for k in data.keys() if k not in FIELD_NAMES[template])  # F6 字典序
    meta = {
        "template": template,
        "title": title,
        "filled_fields": filled,
        "missing_fields": missing,
        "fields": details,
        "summary": {
            "total": len(details),
            "filled": len(filled),
            "missing": len(missing),
            "missing_required": len(missing_required),
            "missing_required_names": missing_required,
            "unknown_keys": unknown,
        },
        "outputs": {"docx": str(docx_path), "fields_json": str(fields_path)},  # S4
    }
    return meta, values


# ---------------------------------------------------------------- docx 生成（V1–V8）

def set_page(section) -> None:
    """V1：A4 + GB/T 9704 版心（上 3.7 / 下 3.5 / 左 2.8 / 右 2.6 cm）。"""
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)


def style_run(run, east: str, size: Pt) -> None:
    """run 字体声明：eastAsia=east、ascii/hAnsi=（正文 TNR，其余与 east 同）、字号。"""
    run.font.size = size
    run.font.name = ASCII_FONT if east == BODY_FONT else east
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)


def set_para_base(p, indent: bool) -> None:
    """段落基础版式：固定行距 28 磅、段前段后 0；indent=True 时首行缩进两字符
    （w:ind/@w:firstLineChars=200，V4/V6）。"""
    pf = p.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if indent:
        p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200")


def add_title(doc, text: str) -> None:
    """V2：首段水平居中，黑体二号。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(text), TITLE_FONT, TITLE_SIZE)


def add_body(doc, text: str, indent: bool = True) -> None:
    """V4：正文段，仿宋三号 / TNR / 28 磅固定行距 / 首行缩进两字符。"""
    p = doc.add_paragraph()
    set_para_base(p, indent)
    if text:
        style_run(p.add_run(text), BODY_FONT, BODY_SIZE)


def add_section_header(doc, text: str) -> None:
    """V6：小节节头，黑体、首行缩进两字符。"""
    p = doc.add_paragraph()
    set_para_base(p, True)
    style_run(p.add_run(text), TITLE_FONT, BODY_SIZE)


def add_signature(doc, text: str) -> None:
    """V5：落款行，右对齐 + 右缩进两字（≈32pt）。text 为空时仍产出空段（不破坏行序）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf.right_indent = SIGN_RIGHT_INDENT
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if text:
        style_run(p.add_run(text), BODY_FONT, BODY_SIZE)


def emit_list(doc, values: dict, name: str) -> None:
    """V6/F4：条目以 1．2．3．（全角点）编号、首行缩进两字符；
    必填列表缺失 → 单个占位条目 1．____。"""
    items = values.get(name)
    if not items:
        items = [PLACEHOLDER]
    for i, item in enumerate(items, 1):
        add_body(doc, "%d．%s" % (i, item))


def contact_line(values: dict, req_contact: bool, req_phone: bool) -> str:
    """联系人/联系电话行内标签（F2：选填缺失标签保留、值留空）。"""
    return "联系人：%s%s联系电话：%s" % (
        pick(values, "联系人", req_contact), CONTACT_SEP, pick(values, "联系电话", req_phone))


def build_doc(template: str, title: str, values: dict) -> Document:
    """按模板公式化成文（V2–V6）。段落顺序：标题 → 称谓/基本信息 → 正文/小节 → 落款。"""
    doc = Document()
    set_page(doc.sections[0])
    add_title(doc, title)

    if template == "周报":
        add_body(doc, "填报人：%s%s周期：%s" % (
            pick(values, "填报人", True), CONTACT_SEP, pick(values, "周期", True)))
        add_section_header(doc, "%s、本周工作内容" % CN_NUMS[0])
        emit_list(doc, values, "本周工作内容")
        add_section_header(doc, "%s、下周工作计划" % CN_NUMS[1])
        emit_list(doc, values, "下周工作计划")
        if "问题与需协调事项" in values:              # F4：选填小节缺失整节省略
            add_section_header(doc, "%s、问题与需协调事项" % CN_NUMS[2])
            emit_list(doc, values, "问题与需协调事项")
        add_signature(doc, "报送日期：%s" % pick(values, "报送日期", False))

    elif template == "请示函":
        add_body(doc, "%s：" % pick(values, "主送机关", True), indent=False)   # V3 称谓顶格
        add_body(doc, pick(values, "请示缘由", True))
        add_body(doc, pick(values, "请示事项", True))
        add_body(doc, contact_line(values, False, False))
        add_signature(doc, pick(values, "请示单位", True))
        add_signature(doc, pick(values, "成文日期", False))

    elif template == "会议通知":
        add_body(doc, "%s：" % pick(values, "主送对象", True), indent=False)   # V3 称谓顶格
        add_body(doc, "会议时间：%s" % pick(values, "会议时间", True))
        add_body(doc, "会议地点：%s" % pick(values, "会议地点", True))
        add_body(doc, "参会人员：%s" % pick(values, "参会人员", True))
        add_body(doc, "会议议题：%s" % pick(values, "会议议题", False))
        if "会议要求" in values:                      # F4：选填小节缺失整节省略
            add_section_header(doc, "%s、会议要求" % CN_NUMS[0])
            emit_list(doc, values, "会议要求")
        add_body(doc, contact_line(values, False, False))
        add_signature(doc, pick(values, "召开单位", True))
        add_signature(doc, pick(values, "发文日期", False))

    else:  # 工作总结
        add_section_header(doc, "%s、工作回顾" % CN_NUMS[0])
        emit_list(doc, values, "工作回顾")
        add_section_header(doc, "%s、主要成绩" % CN_NUMS[1])
        emit_list(doc, values, "主要成绩")
        if "存在问题" in values:                      # F4：选填小节缺失整节省略
            add_section_header(doc, "%s、存在问题" % CN_NUMS[2])
            emit_list(doc, values, "存在问题")
        add_section_header(doc, "%s、下一步工作打算" % CN_NUMS[3])
        emit_list(doc, values, "下一步工作打算")
        add_signature(doc, pick(values, "总结主体", True))
        add_signature(doc, pick(values, "成文日期", False))

    return doc


# ---------------------------------------------------------------- CLI（C1–C3）

def fail(msg: str) -> int:
    """C2：stderr 中文错误、退出码 2；调用前未触碰文件系统，零产物残留。"""
    print("错误：%s" % msg, file=sys.stderr)
    return 2


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="四类中文事务文书生成引擎（周报/请示函/会议通知/工作总结）")
    ap.add_argument("--template", required=True, choices=list(TEMPLATES),
                    help="文书类型（四选一）")
    ap.add_argument("--data", required=True, help="数据 JSON 文件路径（UTF-8，顶层为对象）")
    ap.add_argument("--outdir", required=True, help="输出目录（不存在则递归创建）")
    args = ap.parse_args(argv)

    # ---- 先完成全部输入校验（C2），再触碰输出文件系统 ----
    data_path = Path(args.data)
    if not data_path.is_file():
        return fail("数据文件不存在：%s" % args.data)
    try:
        text = data_path.read_text(encoding="utf-8-sig")
    except OSError as e:
        return fail("数据文件不可读：%s（%s）" % (args.data, e))
    except UnicodeDecodeError as e:
        return fail("数据文件不是 UTF-8 编码：%s（%s）" % (args.data, e))
    try:
        data = json.loads(text)
    except ValueError as e:
        return fail("数据文件不是合法 JSON：%s（%s）" % (args.data, e))
    if not isinstance(data, dict):
        return fail("数据 JSON 顶层必须是对象（键值对），实际为 %s"
                    % type(data).__name__)

    # ---- C1：生成产物 ----
    outdir = Path(args.outdir)
    docx_path = outdir / DOCX_NAME
    fields_path = outdir / FIELDS_NAME

    template = args.template
    details, _ = build_fields(template, data)
    title = build_title(template, {d["name"]: d["value"] for d in details
                                   if d["status"] == "filled"})
    meta, values = build_meta(template, data, title, docx_path, fields_path)
    doc = build_doc(template, title, values)

    try:
        outdir.mkdir(parents=True, exist_ok=True)
        doc.save(str(docx_path))
        fields_path.write_text(
            json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except OSError as e:
        return fail("写入产物失败：%s（%s）" % (args.outdir, e))

    s = meta["summary"]
    print("已生成：%s（模板=%s，标题=%s，字段 filled/total=%d/%d，"
          "必填缺失=%s，模板外键=%s）"
          % (docx_path, template, title, s["filled"], s["total"],
             s["missing_required_names"] or "无", s["unknown_keys"] or "无"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
