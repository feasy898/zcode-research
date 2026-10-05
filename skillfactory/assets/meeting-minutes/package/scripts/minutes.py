#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""minutes.py — meeting-minutes 技能包契约入口

把会议口语转写稿（txt，每行一条 `[hh:mm] 说话人: 内容`）确定性地转成三段式纪要：

    python minutes.py --input <转写.txt> --outdir <输出目录>

产物（固定文件名，写死在 <outdir>/ 下）：
    纪要.docx    一、决议事项（编号条目）／二、待办事项（3 列表格）／三、风险与关注（编号条目）
    summary.json 四键条数统计：决议事项 / 待办事项 / 风险与关注 / 合计

行为规格见资产根 spec.md（R1–R7）。仅依赖 python-docx 与标准库。
"""

import argparse
import json
import os
import re
import sys

# R1.1 发言行解析（时间括号 1–2 位时:2 位分；说话人不含冒号；中英文冒号均可）
LINE_RE = re.compile(r"^\s*\[(\d{1,2}:\d{2})\]\s*([^：:]+?)\s*[：:]\s*(.+?)\s*$")

# R2 关键词分类表：一行至多归一类，按优先级 决议 > 待办 > 风险 取第一个命中
CATEGORY_KEYWORDS = (
    ("决议事项", ("决定", "议定", "同意", "结论")),
    ("待办事项", ("需要", "负责", "跟进", "截止")),
    ("风险与关注", ("风险", "延期", "不确定")),
)

# 人名字符集（R4）：中英文、数字、间隔点，1–6 字
NAME_CHARS = r"[0-9A-Za-z\u4e00-\u9fff·]"
# R4.3/R4.4 要求捕获名后跟标点/空白/行尾
NAME_BOUNDARY = r"(?=[，。；、！？!?.,;:：\s]|$)"

# R4 负责人启发式，按序尝试；'capture_dealias' 表示捕获组为 我/我们/自己 时替换为说话人
OWNER_RULES = (
    ("speaker", re.compile(r"我(?:们)?负责")),
    ("capture", re.compile(r"(?:由|让|交给)(" + NAME_CHARS + r"{1,6}?)负责")),
    ("capture_dealias",
     re.compile(r"负责人\s*(?:是|：|:)\s*(" + NAME_CHARS + r"{1,6}?)" + NAME_BOUNDARY)),
    ("capture", re.compile(r"@(" + NAME_CHARS + r"{1,6}?)" + NAME_BOUNDARY)),
)

# R5 期限启发式，按序尝试（顺序即优先级）；捕获组拼接即为提取结果
DEADLINE_RULES = (
    # 1 完整日期：yyyy年m月d日 / yyyy-m-d / yyyy/m/d 等变体（可带 之前/前）
    re.compile(r"(\d{4}\s*[年/\-.]\s*\d{1,2}\s*[月/\-.]\s*\d{1,2}\s*日?)(之前|前)?"),
    # 2 m月d日（可带 之前/前）
    re.compile(r"(\d{1,2}月\d{1,2}日)(之前|前)?"),
    # 3 (本周|下周|这周)?周X/星期X（前缀只参与命中、不进提取结果；后缀可带 之前/前/以前/底）
    re.compile(r"(?:本周|下周|这周)?((?:周|星期)[一二三四五六日天])(之前|前|以前|底)?"),
    # 4 本(周|月|季度)(底|末|内)?
    re.compile(r"(本(?:周|月|季度)(?:底|末|内)?)"),
    # 5 今天/明天/后天/月底前?/下月/本周末
    re.compile(r"(本周末|今天|明天|后天|月底前|月底|下月)"),
    # 6 截止(于|到|至)?X：X 取 1–8 个非空白且不属于 ，。；; 的字符
    re.compile(r"截止(?:于|到|至)?([^\s，。；;]{1,8})"),
)

UNKNOWN = "待定"
DOCX_NAME = "纪要.docx"
SUMMARY_NAME = "summary.json"
SECTION_DECISION = "一、决议事项"
SECTION_TODO = "二、待办事项"
SECTION_RISK = "三、风险与关注"
TODO_HEADER = ("事项", "负责人", "期限")


def parse_transcript(text):
    """R1：把转写全文解析为 [(hh:mm, 说话人, 内容)]；空行与不合格式行跳过。"""
    utterances = []
    for line in text.splitlines():
        if not line.strip():
            continue  # R1.2 空行忽略
        m = LINE_RE.match(line)
        if not m:
            continue  # R1.3 噪声行跳过
        utterances.append((m.group(1), m.group(2).strip(), m.group(3)))
    return utterances


def classify(content):
    """R2：按子串包含匹配关键词，返回第一个命中的类名（决议>待办>风险），无命中返回 None。"""
    for category, keywords in CATEGORY_KEYWORDS:
        if any(k in content for k in keywords):
            return category
    return None  # R1.4 无关键词的发言不进入任何章节


def extract_owner(content, speaker):
    """R4：从条目内容原文提取负责人；全部失败填 待定。"""
    for kind, pattern in OWNER_RULES:
        m = pattern.search(content)
        if not m:
            continue
        if kind == "speaker":
            return speaker
        name = m.group(1)
        if kind == "capture_dealias" and name in ("我", "我们", "自己"):
            return speaker
        return name
    return UNKNOWN


def extract_deadline(content):
    """R5：从条目内容原文按序提取期限；全部失败填 待定。"""
    for pattern in DEADLINE_RULES:
        m = pattern.search(content)
        if m:
            return "".join(g for g in m.groups() if g)
    return UNKNOWN


def _style_run(run, size=None, bold=False):
    """R3.7：中文显式宋体（含 eastAsia）。"""
    if size is not None:
        run.font.size = size
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "宋体")


def build_docx(source_name, buckets, out_path):
    """R3：生成三段式纪要 docx。buckets = {类名: [(hh:mm, 说话人, 内容)]}（保持原行序）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    counts = {k: len(v) for k, v in buckets.items()}
    total = sum(counts.values())

    doc = Document()

    # R3.1 标题：居中 18pt 加粗
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(title.add_run("会议纪要"), size=Pt(18), bold=True)

    # R3.2 元信息行
    meta = doc.add_paragraph()
    _style_run(meta.add_run(
        "来源：%s ｜ 解析发言 %d 条（决议 %d · 待办 %d · 风险 %d）"
        % (source_name, total, counts["决议事项"], counts["待办事项"],
           counts["风险与关注"])))

    def add_entries(key):
        """R3.4 决议/风险条目：`<i>. 【hh:mm · 说话人】内容原文`，i 从 1 连续。"""
        for i, (when, speaker, content) in enumerate(buckets[key], start=1):
            entry = doc.add_paragraph()
            _style_run(entry.add_run("%d. 【%s · %s】%s" % (i, when, speaker, content)))

    def add_todo_table():
        """R3.5 待办表格：3 列（事项/负责人/期限），每条待办一行。"""
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, TODO_HEADER):
            _style_run(cell.paragraphs[0].add_run(text), bold=True)
        for when, speaker, content in buckets["待办事项"]:
            row = table.add_row()
            _style_run(row.cells[0].paragraphs[0].add_run(content))
            _style_run(row.cells[1].paragraphs[0].add_run(extract_owner(content, speaker)))
            _style_run(row.cells[2].paragraphs[0].add_run(extract_deadline(content)))

    # R3.3 节顺序固定：一、决议事项 → 二、待办事项（表格）→ 三、风险与关注
    head = doc.add_paragraph()
    _style_run(head.add_run(SECTION_DECISION), size=Pt(14), bold=True)
    add_entries("决议事项")

    head = doc.add_paragraph()
    _style_run(head.add_run(SECTION_TODO), size=Pt(14), bold=True)
    add_todo_table()

    head = doc.add_paragraph()
    _style_run(head.add_run(SECTION_RISK), size=Pt(14), bold=True)
    add_entries("风险与关注")

    doc.save(out_path)
    return counts, total


def build_summary(counts, out_path):
    """R6：四键计数 summary.json（UTF-8、缩进 2、不转义中文）。"""
    summary = {
        "决议事项": counts["决议事项"],
        "待办事项": counts["待办事项"],
        "风险与关注": counts["风险与关注"],
        "合计": counts["决议事项"] + counts["待办事项"] + counts["风险与关注"],
    }
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    return summary


def main(argv=None):
    argv = sys.argv if argv is None else argv
    parser = argparse.ArgumentParser(
        prog="minutes.py",
        description="会议转写稿 → 三段式纪要（纪要.docx + summary.json）")
    parser.add_argument("--input", required=True, help="转写 txt 路径")
    parser.add_argument("--outdir", required=True, help="输出目录（不存在自动创建）")
    args = parser.parse_args(argv[1:])

    # R7.1 输入不存在：非 0 退出、报错、不产出文件（先查输入，后建目录）
    if not os.path.isfile(args.input):
        print("错误：输入文件不存在: %s" % args.input, file=sys.stderr)
        return 2
    try:
        with open(args.input, "r", encoding="utf-8-sig") as f:  # R1.5 容忍 BOM
            text = f.read()
    except OSError as exc:
        print("错误：输入文件无法读取: %s" % exc, file=sys.stderr)
        return 2

    buckets = {key: [] for key, _kw in CATEGORY_KEYWORDS}
    for when, speaker, content in parse_transcript(text):
        category = classify(content)
        if category is not None:
            buckets[category].append((when, speaker, content))

    try:
        os.makedirs(args.outdir, exist_ok=True)
        docx_path = os.path.join(args.outdir, DOCX_NAME)
        summary_path = os.path.join(args.outdir, SUMMARY_NAME)
        counts, total = build_docx(os.path.basename(args.input), buckets, docx_path)
        summary = build_summary(counts, summary_path)
    except OSError as exc:
        print("错误：产物写入失败: %s" % exc, file=sys.stderr)
        return 2

    print("OK %s | %s | 决议 %d · 待办 %d · 风险 %d · 合计 %d"
          % (docx_path, summary_path, summary["决议事项"], summary["待办事项"],
             summary["风险与关注"], summary["合计"]))
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001 - 非 tty 场景可能不支持
        pass
    sys.exit(main())
