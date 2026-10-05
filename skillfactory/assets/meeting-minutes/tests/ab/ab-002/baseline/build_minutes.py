# -*- coding: utf-8 -*-
"""
meeting-minutes baseline builder — case3 (HIS 系统升级会)

三段式会议纪要：一、决议事项；二、待办事项；三、风险提示。
数据单源（本文件内的 DECISIONS/ACTIONS/RISKS），同时产出 纪要.docx / summary.json / content.md，
保证三者一致。

归类规则（依据任务 rubric）：
  1) 一行至多归入一类；优先级 决议 > 待办 > 风险。
     [08:43] 同时含“延期”（风险词）与“结论”（决议词）→ 归决议事项，不入风险。
  2) 待办提取：负责人 / 期限；抓不到的如实填“待定”。
     [08:38] 负责人=医务处、期限=周五（取“截止到周五”，不得误判为“明天”）。
  3) 噪声排除：寒暄（08:30）、口误更正（08:41，含“记错了”）、散会（08:45）；
     08:32/08:33 为讨论性发言，最终以 08:35 决议为准，不单列条目。
  4) 每条目附“原文出处”，逐字可在 case3.txt 中找到。
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(r"D:\workspace\zcode研究")
TRANSCRIPT = ROOT / "skillfactory" / "assets" / "meeting-minutes" / "oracle" / "inputs" / "case3.txt"
OUTDIR = ROOT / "skillfactory" / "assets" / "meeting-minutes" / "tests" / "ab" / "ab-002" / "baseline"

# ---------------------------------------------------------------- 数据（单源）
MEETING = {
    "title": "会议纪要",
    "topic": "HIS 系统升级",
    "topic_source": "[08:30] 张科长: 各位早，人都齐了就讲正事，今天只议HIS系统升级这一件事。",
    "time_span": "08:30–08:45（转写稿未注明日期）",
    "attendees": "张科长、刘护士长、李主任、厂商小吴（依发言人整理）",
}

DECISIONS = [
    {
        "id": "决议 1",
        "content": "停机窗口调整为零点到四点，会议通知由张科长负责发布。",
        "source": "[08:35] 张科长: 同意，窗口改到零点到四点，通知我来发。",
    },
    {
        "id": "决议 2",
        "content": "升级当天一旦出现延期迹象，立即回滚。",
        "source": "[08:43] 张科长: 那就好。升级当天一旦出现延期迹象，立即回滚，这个结论请各位记牢。",
    },
]

ACTIONS = [
    {
        "id": "1",
        "task": "发布停机公告",
        "owner": "医务处",
        "deadline": "部署当晚",
        "source": "[08:36] 厂商小吴: 部署当晚需要医务处发停机公告，各科室还要演练一遍纸质病历流程。",
    },
    {
        "id": "2",
        "task": "演练一遍纸质病历流程",
        "owner": "各科室",
        "deadline": "待定",
        "source": "[08:36] 厂商小吴: 部署当晚需要医务处发停机公告，各科室还要演练一遍纸质病历流程。",
    },
    {
        "id": "3",
        "task": "发布停机公告（明天发OA）；各科室把应急预案报上来",
        "owner": "医务处",
        "deadline": "周五",
        "source": "[08:38] 李主任: 公告由医务处负责，明天发OA，截止到周五各科室把应急预案报上来。",
    },
]

RISKS = [
    {
        "id": "风险 1",
        "content": "老库迁移脚本厂商还没给，存在不确定性。",
        "source": "[08:40] 张科长: 有个不确定的点，老库迁移脚本厂商还没给，这是个风险。",
    },
]

EXCLUDED = [
    {"time": "08:30", "type": "寒暄/开场", "line": "[08:30] 张科长: 各位早，人都齐了就讲正事，今天只议HIS系统升级这一件事。"},
    {"time": "08:32", "type": "讨论性发言（初步方案，被 08:35 决议取代），不单列条目", "line": "[08:32] 厂商小吴: 新版本下周凌晨部署，停机窗口初步定在凌晨一点到五点。"},
    {"time": "08:33", "type": "讨论性发言（异议，促成 08:35 决议），不单列条目", "line": "[08:33] 刘护士长: 这个窗口不行，夜班五点半交班，系统必须提前恢复，需要留足缓冲。"},
    {"time": "08:41", "type": "口误更正（含“记错了”），按噪声排除", "line": "[08:41] 厂商小吴: 不好意思我记错了，脚本上周就发过邮件了，我等下再转发一次。"},
    {"time": "08:45", "type": "散会行，按噪声排除", "line": "[08:45] 张科长: 谢谢各位，散会。"},
]

CLASSIFICATION = {
    "one_line_one_class": "一行至多归入一类",
    "priority": "决议 > 待办 > 风险",
    "priority_example": "[08:43] 同时含“延期”（风险词）与“结论”（决议词）→ 按优先级归入决议事项（决议 2），不入风险提示",
    "action_heuristics": "负责人/期限抓不到的行如实填“待定”；[08:38] 负责人=医务处、期限=周五（取“截止到周五”，不误判为“明天”）",
    "noise_excluded": ["08:30 寒暄", "08:41 口误更正（记错了）", "08:45 散会", "08:32/08:33 讨论性发言（以 08:35 决议为准，不单列）"],
}


def _set_east_asian(style, name="微软雅黑"):
    try:
        style.font.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), name)
    except Exception:
        pass


def _add_source_para(doc, text):
    p = doc.add_paragraph()
    run = p.add_run("原文出处：" + text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    return p


def build_docx(path):
    doc = Document()
    _set_east_asian(doc.styles["Normal"])

    h = doc.add_heading(MEETING["title"], level=0)
    p = doc.add_paragraph()
    r = p.add_run("HIS 系统升级会")
    r.bold = True
    r.font.size = Pt(14)

    info = doc.add_table(rows=4, cols=2)
    info.style = "Table Grid"
    info_rows = [
        ("会议主题", MEETING["topic"] + "（来源 [08:30]“今天只议HIS系统升级这一件事”）"),
        ("记录时间范围", MEETING["time_span"]),
        ("与会人员", MEETING["attendees"]),
        ("条目合计", "决议事项 2 条 · 待办事项 3 条 · 风险提示 1 条 · 共 6 条"),
    ]
    for i, (k, v) in enumerate(info_rows):
        info.rows[i].cells[0].text = k
        info.rows[i].cells[1].text = v
        for run in info.rows[i].cells[0].paragraphs[0].runs:
            run.bold = True

    # 一、决议事项
    doc.add_heading("一、决议事项", level=1)
    for item in DECISIONS:
        p = doc.add_paragraph()
        run = p.add_run(item["id"] + "：" + item["content"])
        run.bold = True
        _add_source_para(doc, item["source"])

    # 二、待办事项
    doc.add_heading("二、待办事项", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    header = ("序号", "事项", "负责人", "期限", "原文出处（case3.txt）")
    for j, text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for a in ACTIONS:
        row = table.add_row()
        row.cells[0].text = a["id"]
        row.cells[1].text = a["task"]
        row.cells[2].text = a["owner"]
        row.cells[3].text = a["deadline"]
        row.cells[4].text = a["source"]

    # 三、风险提示
    doc.add_heading("三、风险提示", level=1)
    for item in RISKS:
        p = doc.add_paragraph()
        run = p.add_run(item["id"] + "：" + item["content"])
        run.bold = True
        _add_source_para(doc, item["source"])

    # 附注（非条目）
    p = doc.add_paragraph()
    r = p.add_run("── 附注（整理说明，非纪要条目）──")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    rules = doc.add_paragraph()
    r = rules.add_run(
        "归类规则：一行至多归入一类，优先级 决议 > 待办 > 风险；"
        "[08:43] 同时含“延期”与“结论”，按优先级归入决议事项（决议 2），未计入风险。"
        "待办负责人/期限无法从原文取得的如实填“待定”；[08:38] 期限取“截止到周五”，不取“明天发OA”的“明天”。"
    )
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    for e in EXCLUDED:
        p = doc.add_paragraph()
        r = p.add_run(f"未纳入条目 [{e['time']}]（{e['type']}）：{e['line']}")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(path)


def build_summary(path):
    counts = {
        "决议事项": len(DECISIONS),
        "待办事项": len(ACTIONS),
        "风险提示": len(RISKS),
    }
    counts["total"] = sum(counts[k] for k in ("决议事项", "待办事项", "风险提示"))
    summary = {
        "task": "meeting-minutes",
        "case": "case3",
        "source": "skillfactory/assets/meeting-minutes/oracle/inputs/case3.txt",
        "meeting_topic": MEETING["topic"],
        "format": "三段式：决议事项 / 待办事项 / 风险提示",
        "counts": counts,
        "decisions": DECISIONS,
        "action_items": ACTIONS,
        "risks": RISKS,
        "classification": CLASSIFICATION,
        "excluded_noise": EXCLUDED,
    }
    path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def build_content_md(path, summary):
    lines = []
    lines.append("# 会议纪要（HIS 系统升级会）— 产物说明\n")
    lines.append("产物：`纪要.docx`（三段式会议纪要）、`summary.json`（结构化计数与条目，与文档一致）、本说明文件。\n")
    lines.append("来源转写稿：`skillfactory/assets/meeting-minutes/oracle/inputs/case3.txt`（10 行发言记录，08:30–08:45）。\n")

    lines.append("## 完整文本内容（与 纪要.docx 一致）\n")
    lines.append(f"**{MEETING['title']}：HIS 系统升级会**\n")
    lines.append("| 项目 | 内容 |")
    lines.append("|---|---|")
    lines.append(f"| 会议主题 | {MEETING['topic']}（来源 [08:30]“今天只议HIS系统升级这一件事”） |")
    lines.append(f"| 记录时间范围 | {MEETING['time_span']} |")
    lines.append(f"| 与会人员 | {MEETING['attendees']} |")
    lines.append("| 条目合计 | 决议事项 2 条 · 待办事项 3 条 · 风险提示 1 条 · 共 6 条 |\n")

    lines.append("### 一、决议事项（2 条）\n")
    for item in DECISIONS:
        lines.append(f"- **{item['id']}：{item['content']}**")
        lines.append(f"  - 原文出处：{item['source']}")
    lines.append("")

    lines.append("### 二、待办事项（3 行数据，含 负责人/期限 列）\n")
    lines.append("| 序号 | 事项 | 负责人 | 期限 | 原文出处（case3.txt） |")
    lines.append("|---|---|---|---|---|")
    for a in ACTIONS:
        lines.append(f"| {a['id']} | {a['task']} | {a['owner']} | {a['deadline']} | {a['source']} |")
    lines.append("")

    lines.append("### 三、风险提示（1 条）\n")
    for item in RISKS:
        lines.append(f"- **{item['id']}：{item['content']}**")
        lines.append(f"  - 原文出处：{item['source']}")
    lines.append("")

    lines.append("── 附注（整理说明，非纪要条目）──\n")
    lines.append(
        "归类规则：一行至多归入一类，优先级 决议 > 待办 > 风险；[08:43] 同时含“延期”与“结论”，"
        "按优先级归入决议事项（决议 2），未计入风险。\n"
    )
    for e in EXCLUDED:
        lines.append(f"- 未纳入条目 [{e['time']}]（{e['type']}）：{e['line']}")
    lines.append("")

    lines.append("## 结构说明与归类依据\n")
    lines.append("1. **三段式结构**：一、决议事项；二、待办事项（表格，含 负责人/期限 列，3 行数据）；三、风险提示。")
    lines.append("2. **一行至多归一类，优先级 决议 > 待办 > 风险**：[08:43] 行（张科长）同时含“延期”（风险词）与“结论”（决议词），归入决议事项（决议 2），不计入风险。")
    lines.append("3. **决议 2 条**：[08:35]（“同意”→停机窗口改零点到四点，通知张科长发）、[08:43]（延期迹象立即回滚）。[08:35] 虽含待办性质的“通知我来发”，按 决议>待办 整行归决议。")
    lines.append("4. **待办 3 条**：[08:36] 拆为两条（① 医务处发停机公告，期限“部署当晚”；② 各科室演练纸质病历流程，原文未给该句单独期限，如实填“待定”）；[08:38] 整行一条（负责人=医务处，期限=周五——取“截止到周五”，未被误判为“明天发OA”的“明天”；“明天发OA”与各科室报应急预案保留在事项描述与原文出处中）。")
    lines.append("5. **风险 1 条**：[08:40] 老库迁移脚本厂商还没给。")
    lines.append("6. **噪声抑制**：[08:30] 寒暄、[08:41] 口误更正（“记错了”）、[08:45] 散会行均未混入三个条目节；[08:32]/[08:33] 为讨论性发言（初步方案与异议），最终以 [08:35] 决议为准，不单列。")
    lines.append("7. **可溯源**：每条目附“原文出处”，逐字可在 case3.txt 中找到（验证脚本 verify_minutes.py 已核对子串匹配）。")
    lines.append("8. **计数一致性**：决议 2 + 待办 3 + 风险 1 = 6，与 纪要.docx 及 summary.json 一致。\n")

    path.write_text("\n".join(lines), encoding="utf-8")


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)
    docx_path = OUTDIR / "纪要.docx"
    build_docx(docx_path)
    summary = build_summary(OUTDIR / "summary.json")
    build_content_md(OUTDIR / "content.md", summary)
    print("written:", docx_path)
    print("counts:", summary["counts"])


if __name__ == "__main__":
    main()
