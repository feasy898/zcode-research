# -*- coding: utf-8 -*-
"""verify_minutes.py — 独立核对 纪要.docx / summary.json 是否满足 rubric。"""
import json
import zipfile
from pathlib import Path

from docx import Document

BASE = Path(r"D:\workspace\zcode研究\skillfactory\assets\meeting-minutes\tests\ab\ab-002\baseline")
TRANSCRIPT = Path(r"D:\workspace\zcode研究\skillfactory\assets\meeting-minutes\oracle\inputs\case3.txt")

results = []


def check(name, ok, detail=""):
    results.append((name, ok, detail))
    print(("PASS" if ok else "FAIL"), "-", name, ("| " + detail if detail else ""))


# ---- 载入
transcript = TRANSCRIPT.read_text(encoding="utf-8").replace("\r\n", "\n")
doc = Document(BASE / "纪要.docx")
summary = json.loads((BASE / "summary.json").read_text(encoding="utf-8"))

paras = [p.text for p in doc.paragraphs]
all_text = "\n".join(paras)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            all_text += "\n" + c.text

# ---- zip 完整性
with zipfile.ZipFile(BASE / "纪要.docx") as z:
    bad = z.testzip()
check("docx zip 完整性", bad is None, "zip testzip -> None" if bad is None else f"bad entry {bad}")

# ---- 定位三个条目节（以 Heading 文本切分）
def section_slice(start_marker, end_marker):
    i = next(k for k, t in enumerate(paras) if t.strip().startswith(start_marker))
    j = next(k for k, t in enumerate(paras) if t.strip().startswith(end_marker))
    return "\n".join(paras[i:j])

sec_dec = section_slice("一、决议事项", "二、待办事项")
sec_act_head_i = next(k for k, t in enumerate(paras) if t.strip().startswith("二、待办事项"))
sec_risk_i = next(k for k, t in enumerate(paras) if t.strip().startswith("三、风险提示"))
sec_act = "\n".join(paras[sec_act_head_i:sec_risk_i])
sec_risk_and_after = "\n".join(paras[sec_risk_i:])
sec_risk = sec_risk_and_after.split("── 附注")[0]

# ---- 1) 优先级：08:43（含“延期”又含“结论”）归决议而非风险；一行至多一类
check("优先级：[08:43] 在决议节", "[08:43]" in sec_dec)
check("优先级：[08:43] 不在风险节", "[08:43]" not in sec_risk)
check("优先级：决议节含 [08:35]", "[08:35]" in sec_dec)
check("一行至多一类：决议节不含待办行 [08:36]/[08:38]", "[08:36]" not in sec_dec and "[08:38]" not in sec_dec)
check("一行至多一类：风险节仅 [08:40]", "[08:40]" in sec_risk and "[08:43]" not in sec_risk and "[08:36]" not in sec_risk and "[08:38]" not in sec_risk)

# ---- 2) 条数：决议 2 / 待办 3 / 风险 1，文档与 summary.json 一致，合计 6
check("summary.json counts = 2/3/1",
      {k: summary["counts"][k] for k in ("决议事项", "待办事项", "风险提示")} == {"决议事项": 2, "待办事项": 3, "风险提示": 1},
      f"counts={summary['counts']}")
check("summary.json 合计 = 6", summary["counts"]["total"] == 6)
check("summary 条目数组长度与 counts 一致",
      len(summary["decisions"]) == 2 and len(summary["action_items"]) == 3 and len(summary["risks"]) == 1)
check("文档决议节含 决议 1/2", "决议 1" in sec_dec and "决议 2" in sec_dec and "决议 3" not in sec_dec)
check("文档风险节含 风险 1", "风险 1" in sec_risk and "风险 2" not in sec_risk)

# ---- 3) 待办表格：含 负责人/期限 列，数据行 3 行
actions_table = None
for t in doc.tables:
    hdr = [c.text.strip() for c in t.rows[0].cells]
    if "负责人" in hdr and "期限" in hdr:
        actions_table = t
        hdr_actions = hdr
        break
check("待办表存在且表头含 负责人/期限", actions_table is not None, f"表头={hdr_actions if actions_table else None}")
data_rows = actions_table.rows[1:] if actions_table else []
check("待办表数据行 = 3 行", len(data_rows) == 3, f"实际 {len(data_rows)} 行")

rows = [[c.text.strip() for c in r.cells] for r in data_rows]
i_owner = hdr_actions.index("负责人")
i_deadline = hdr_actions.index("期限")
i_src = hdr_actions.index("原文出处（case3.txt）")

# ---- 4) 待办启发式
row_0838 = next((r for r in rows if "[08:38]" in r[i_src]), None)
check("[08:38] 行 负责人=医务处", row_0838 is not None and row_0838[i_owner] == "医务处",
      f"负责人={row_0838[i_owner] if row_0838 else None}")
check("[08:38] 行 期限=周五（截止到周五）", row_0838 is not None and row_0838[i_deadline] == "周五",
      f"期限={row_0838[i_deadline] if row_0838 else None}")
deadlines = [r[i_deadline] for r in rows]
check("无任何期限被误判为“明天”", all(d != "明天" for d in deadlines), f"期限列={deadlines}")
check("抓不到期限的行如实填“待定”", "待定" in deadlines, f"期限列={deadlines}")

# ---- 5) 噪声抑制：08:30 寒暄 / 08:41 口误更正 / 08:45 散会 不混入条目节
for ts, label in (("08:30", "寒暄"), ("08:41", "口误更正"), ("08:45", "散会")):
    in_items = (f"[{ts}]" in sec_dec) or (f"[{ts}]" in sec_act) or (f"[{ts}]" in sec_risk)
    check(f"噪声 [{ts}] {label} 未混入三个条目节", not in_items)
check("附注中明确列出排除项", ("[08:30]" in sec_risk_and_after and "[08:41]" in sec_risk_and_after and "[08:45]" in sec_risk_and_after))

# ---- 6) 逐字溯源：每条原文出处可在 case3.txt 中逐字找到
srcs = [d["source"] for d in summary["decisions"]] + [r["source"] for r in summary["risks"]] + \
       [a["source"] for a in summary["action_items"]]
missing = [s for s in srcs if s not in transcript]
check("summary 全部原文出处逐字命中 case3.txt", not missing, f"未命中 {len(missing)} 条")
doc_missing = [s for s in srcs if s not in all_text]
check("文档中全部原文出处逐字出现", not doc_missing, f"未命中 {len(doc_missing)} 条")

# ---- 7) summary.json 与文档一致（待办表逐格比对）
table_ok = True
detail = []
for a, r in zip(summary["action_items"], rows):
    pair = [
        (a["task"], r[hdr_actions.index("事项")]),
        (a["owner"], r[i_owner]),
        (a["deadline"], r[i_deadline]),
        (a["source"], r[i_src]),
    ]
    for x, y in pair:
        if x != y:
            table_ok = False
            detail.append(f"{x!r} != {y!r}")
check("summary.action_items 与待办表逐格一致", table_ok, "; ".join(detail))

print()
n_pass = sum(1 for _, ok, _ in results if ok)
print(f"== {n_pass}/{len(results)} checks passed ==")
