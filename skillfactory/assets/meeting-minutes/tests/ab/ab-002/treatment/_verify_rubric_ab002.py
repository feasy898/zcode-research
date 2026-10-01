# -*- coding: utf-8 -*-
"""_verify_rubric_ab002.py — ab-002 treatment 验收脚本（中间产物，留在 outdir）

1) 调技能自带 eval/runner.py 的 check_case("case3", staging, oracle/out)：5 项机检。
2) rubric 专项：08:43 优先级归决议、条数 2/3/1=6、待办表 3 行及负责人/期限提取、
   08:30/08:41/08:45 噪声零泄漏、全部条目逐字溯源、与 oracle 金标准逐段逐格一致。

用法：python _verify_rubric_ab002.py   （全过 exit 0，任一失败 exit 1）
"""
import json
import os
import shutil
import sys

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ASSET = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))
sys.path.insert(0, os.path.join(ASSET, "eval"))
import runner  # noqa: E402  skillfactory assets meeting-minutes eval/runner.py

CASE = "case3"
OUT = HERE
STAGING = os.path.join(OUT, "_staging")
ORACLE = os.path.join(ASSET, "oracle", "out")
INPUT = os.path.join(ASSET, "oracle", "inputs", CASE + ".txt")
DOCX = os.path.join(OUT, "纪要.docx")
SUMMARY = os.path.join(OUT, "summary.json")

FAILS = []


def check(name, ok, detail):
    print("[%s] %s — %s" % ("PASS" if ok else "FAIL", name, detail))
    if not ok:
        FAILS.append(name)


# ---- 0) staging（runner 要求 <cand>/<case>/ 布局），从当前产物刷新 ----
os.makedirs(os.path.join(STAGING, CASE), exist_ok=True)
shutil.copyfile(DOCX, os.path.join(STAGING, CASE, "纪要.docx"))
shutil.copyfile(SUMMARY, os.path.join(STAGING, CASE, "summary.json"))
print("staging refreshed:", os.path.join(STAGING, CASE))

# ---- 1) 技能自带 runner check_case（5 项：三节标题/待办表格列/summary一致/逐字溯源/条数容差）----
for c in runner.check_case(CASE, STAGING, ORACLE):
    check("runner/" + c["name"], c["pass"], c["detail"])

# ---- 公共材料 ----
doc = Document(DOCX)
paras = [p.text for p in doc.paragraphs]
tables = doc.tables
all_cells = [cell.text for t in tables for row in t.rows for cell in row.cells]
sections = runner.section_entries(doc)
with open(INPUT, encoding="utf-8-sig") as f:
    lines = f.read().splitlines()
LINE = {ln[1:6]: ln for ln in lines}  # "08:33" -> 原行

# ---- 2) rubric-优先级：08:43（含「延期」又含「结论」）归决议而非风险；一行至多一类 ----
dec = sections["一、决议事项"]
risk = sections["三、风险与关注"]
L0843 = LINE["08:43"].split(": ", 1)[1]
L0835 = LINE["08:35"].split(": ", 1)[1]
L0840 = LINE["08:40"].split(": ", 1)[1]
check("priority/0843_in_decision", len(dec) == 2 and L0843 in dec[1] and "延期" in dec[1]
      and "结论" in dec[1],
      "决议第2条=08:43原句（含延期+结论）: %s" % (dec[1] if len(dec) > 1 else dec))
check("priority/0843_not_in_risk", len(risk) == 1 and L0843 not in risk[0],
      "风险仅1条且不含08:43行: %s" % risk)
check("priority/one_class_per_line",
      L0835 in dec[0] and L0835 not in "".join(all_cells) and L0840 in risk[0]
      and L0840 not in "".join(dec),
      "08:35 只在决议、08:40 只在风险、待办表格无重复行")

# ---- 3) rubric-条数：2/3/1，summary 与文档一致，合计=6 ----
with open(SUMMARY, encoding="utf-8-sig") as f:
    summary = json.load(f)
todo_table = next(t for t in tables
                  if [c.text.strip() for c in t.rows[0].cells] == ["事项", "负责人", "期限"])
data_rows = todo_table.rows[1:]
actual = {"决议事项": len(dec), "待办事项": len(data_rows), "风险与关注": len(risk)}
check("counts/summary_matches_docx",
      summary == {"决议事项": 2, "待办事项": 3, "风险与关注": 1, "合计": 6}
      and summary["决议事项"] == actual["决议事项"]
      and summary["待办事项"] == actual["待办事项"]
      and summary["风险与关注"] == actual["风险与关注"]
      and summary["合计"] == sum(actual.values()) == 6,
      "summary=%s 与 docx 实际 %s 一致，合计=6" % (json.dumps(summary, ensure_ascii=False), actual))
check("counts/meta_line",
      paras[1] == "来源：case3.txt ｜ 解析发言 6 条（决议 2 · 待办 3 · 风险 1）",
      "元信息行: %s" % paras[1])

# ---- 4) rubric-待办启发式：逐行核对 负责人/期限 ----
expect = [
    ("08:33", "待定", "待定"),
    ("08:36", "待定", "待定"),
    ("08:38", "医务处", "周五"),
]
ok_rows, details = True, []
for row, (tmin, owner, deadline) in zip(data_rows, expect):
    content = row.cells[0].text.strip()
    src = LINE[tmin]
    ok = (content in src and "[%s]" % tmin in src
          and row.cells[1].text.strip() == owner and row.cells[2].text.strip() == deadline
          and (owner == "待定" or owner in content) and (deadline == "待定" or deadline in content))
    ok_rows &= ok
    details.append("%s行: 负责人=%s 期限=%s 事项逐字命中=%s"
                   % (tmin, row.cells[1].text.strip(), row.cells[2].text.strip(), content in src))
check("todo/rows_owner_deadline", ok_rows and len(data_rows) == 3, "; ".join(details))
deadline_col = [r.cells[2].text.strip() for r in data_rows]
check("todo/deadline_not_mingTian", "明天" not in "".join(deadline_col),
      "期限列=%s（规则3「周五」优先于规则5「明天」，未被误判）" % deadline_col)

# ---- 5) rubric-噪声抑制：08:30 寒暄 / 08:41 口误 / 08:45 散会 全文零泄漏 ----
NOISE = {
    "08:30": ["各位早", "人都齐了", "只议HIS"],
    "08:41": ["不好意思我记错了", "上周就发过邮件", "等下再转发一次"],
    "08:45": ["谢谢各位", "散会"],
}
full_text = "\n".join(paras + all_cells)
leaked = [tmin + ":" + s for tmin, subs in NOISE.items() for s in subs if s in full_text]
check("noise/no_leakage", not leaked, "寒暄/口误/散会行特征串零命中%s" % (leaked or ""))

# ---- 6) rubric-逐字溯源：每条目（决议/风险/待办事项列）逐字在 case3.txt 找到出处行 ----
entries = [(runner.parse_entry(t) or (None, None, t)) for t in dec + risk]
entries += [(None, None, r.cells[0].text.strip()) for r in data_rows]
bad = []
for tmin_s, speaker, content in entries:
    hit = next((ln for ln in lines if content in ln), None)
    if hit is None:
        bad.append("无出处:" + content[:20])
    elif tmin_s and ("[%s]" % tmin_s not in hit or speaker not in hit):
        bad.append("时间/说话人不符:" + tmin_s)
check("verbatim/all_entries_traceable", not bad and len(entries) == 6,
      "%d/6 条目逐字溯源通过%s" % (len(entries) - len(bad), bad or ""))

# ---- 7) 与 oracle 金标准逐段逐格一致（脚本确定性输出的旁证）----
gold = Document(os.path.join(ORACLE, CASE, "纪要.docx"))
gold_paras = [p.text for p in gold.paragraphs]
gold_cells = [c.text for t in gold.tables for r in t.rows for c in r.cells]
check("gold/equal_to_oracle",
      paras == gold_paras and all_cells == gold_cells
      and json.load(open(SUMMARY, encoding="utf-8-sig"))
      == json.load(open(os.path.join(ORACLE, CASE, "summary.json"), encoding="utf-8-sig")),
      "docx 全部段落(%d)+单元格(%d) 与 oracle/out/case3 逐字一致，summary 一致"
      % (len(paras), len(all_cells)))

print("\n== %d/%d checks passed ==" % (len([1 for _ in results]) - len(FAILS)
                                       if False else 0, 0) if False else
      "\nFAILS: %s" % (FAILS or "无 — 全部通过"))
sys.exit(1 if FAILS else 0)
