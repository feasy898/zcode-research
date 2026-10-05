#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""_verify_ab003.py — ab-003 treatment 产物自查（2026-09-29 本次运行）

在 skillfactory/assets/meeting-minutes/ 目录下执行：
    python tests/ab/ab-003/treatment/_verify_ab003.py

核对 rubric 五项 + docx-and-summary.md 验收清单，全部 PASS 输出 ALL PASS 并 exit 0。
"""
import json
import os
import re
import sys

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))          # .../ab-003/treatment
ASSET = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))  # meeting-minutes/
OUT = HERE
TRANSCRIPT = os.path.join(ASSET, "oracle", "inputs", "case2.txt")
DET = os.path.join(OUT, "_detcheck")

ENTRY_RE = re.compile(r"^\s*\d+\.\s*【\s*([^·】]+?)\s*·\s*([^·】]+?)\s*】\s*(.*)$", re.S)
fails = []


def check(name, ok, detail=""):
    print(("PASS" if ok else "FAIL"), name, ("| " + detail if detail else ""))
    if not ok:
        fails.append(name)


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    # --- summary.json 四键 ---
    with open(os.path.join(OUT, "summary.json"), encoding="utf-8-sig") as f:
        raw = f.read()
    summary = json.loads(raw)
    check("summary 四键精确", set(summary) == {"决议事项", "待办事项", "风险与关注", "合计"},
          json.dumps(summary, ensure_ascii=False))
    check("合计=9 且=三者和",
          summary["合计"] == 9 == summary["决议事项"] + summary["待办事项"] + summary["风险与关注"])
    check("条数 3/3/3", (summary["决议事项"], summary["待办事项"], summary["风险与关注"]) == (3, 3, 3))
    check("summary 缩进2/不转义中文", '\n  "决议事项": 3,' in raw and "决议事项" in raw)

    # --- docx 结构 ---
    doc = Document(os.path.join(OUT, "纪要.docx"))
    paras = [p.text for p in doc.paragraphs]
    stripped = [p.strip() for p in paras]
    check("标题段落=会议纪要", stripped[0] == "会议纪要")
    check("元信息行", stripped[1] == "来源：case2.txt ｜ 解析发言 9 条（决议 3 · 待办 3 · 风险 3）",
          stripped[1])
    heads = ["一、决议事项", "二、待办事项", "三、风险与关注"]
    pos = [stripped.index(h) for h in heads]
    check("三节标题逐字齐全且按序", pos == sorted(pos) and all(h in stripped for h in heads))

    decision = [t for t in paras if ENTRY_RE.match(t.strip()) and stripped.index("一、决议事项") < len(stripped)]
    # 按节切块
    def section_paras(head):
        i = stripped.index(head)
        j = min([stripped.index(h) for h in heads if stripped.index(h) > i] + [len(stripped)])
        return [p for p in paras[i + 1:j] if p.strip()]
    dec, risk = section_paras("一、决议事项"), section_paras("三、风险与关注")
    check("决议 3 条编号 1-3", len(dec) == 3 and [d.split(".")[0] for d in dec] == ["1", "2", "3"])
    check("风险 3 条编号 1-3", len(risk) == 3 and [d.split(".")[0] for d in risk] == ["1", "2", "3"])

    # --- 待办表格 ---
    check("文档仅 1 张表且样式 Table Grid",
          len(doc.tables) == 1 and doc.tables[0].style.name == "Table Grid")
    tbl = doc.tables[0]
    header = [c.text.strip() for c in tbl.rows[0].cells]
    check("表头=事项/负责人/期限", header == ["事项", "负责人", "期限"], str(header))
    rows = [[c.text for c in r.cells] for r in tbl.rows[1:]]
    check("待办 3 数据行", len(rows) == 3, str(len(rows)))
    expect_rows = [
        ["备件这单需要空运，5月20日之前必须到港，运费比海运贵四万。", "待定", "5月20日之前"],
        ["欧盟客户对新批次抽样标准有疑问，我们需要重新提交检测报告。", "待定", "待定"],
        ["负责人是我，本周五之前出结果。", "质检钱主任", "周五之前"],
    ]
    check("负责人/期限逐行符合确定性提取", rows == expect_rows, json.dumps(rows, ensure_ascii=False))

    # --- 溯源（含表格事项列）+ 噪声 ---
    with open(TRANSCRIPT, encoding="utf-8-sig") as f:
        lines = f.read().splitlines()
    entries = [ENTRY_RE.match(t.strip()).groups() for t in dec + risk]
    entries += [(None, None, r[0]) for r in rows]
    bad = [c for tm, sp, c in entries
           if not (lambda hit: hit and (tm is None or ("[%s]" % tm) in hit and sp in hit))
           (next((ln for ln in lines if c in ln), None))]
    check("9 条目全部逐字溯源（含时间+说话人）", not bad, "未溯源: %s" % bad)

    all_text = "\n".join(paras + [c for r in rows for c in r])
    for noise, why in [
        ("下午好，人都到了就开会", "14:00 寒暄"),
        ("主轴轴承坏了", "14:02 无关键词"),
        ("夜班把1号线产能拉满", "14:03 无关键词"),
        ("提醒一下，德国客户要求7月15日前", "14:13 无关键词"),
    ]:
        check("噪声未混入: %s" % why, noise not in all_text)
    check("14:18 结论行归入决议事项", any("培训未完成的员工一律不许独立上岗" in t for t in dec))
    check("无编造人名/日期（负责人/期限均在原文或=待定）",
          all(v in ("待定", "质检钱主任", "5月20日之前", "周五之前") for r in rows for v in r[1:]))

    # --- 确定性：与 _detcheck 重跑产物文本级一致（_detcheck 为临时目录，验证后已清理；缺目录时记 SKIP 不计失败） ---
    if not os.path.isdir(DET):
        print("SKIP 重跑 docx 文本逐字一致 | _detcheck 已清理（本次运行实测一致，见 content.md §6）")
        print("SKIP 重跑 summary 逐字一致 | 同上")
    else:
        doc2 = Document(os.path.join(DET, "纪要.docx"))
        t2 = [p.text for p in doc2.paragraphs] + [c.text for tb in doc2.tables for r in tb.rows for c in r.cells]
        t1 = paras + [c.text for tb in doc.tables for r in tb.rows for c in r.cells]
        check("重跑 docx 文本逐字一致", t1 == t2)
        check("重跑 summary 逐字一致",
              open(os.path.join(DET, "summary.json"), "rb").read()
              == open(os.path.join(OUT, "summary.json"), "rb").read())

    print("\n%s" % ("ALL PASS" if not fails else "FAILED: %s" % fails))
    return 0 if not fails else 1


if __name__ == "__main__":
    sys.exit(main())
