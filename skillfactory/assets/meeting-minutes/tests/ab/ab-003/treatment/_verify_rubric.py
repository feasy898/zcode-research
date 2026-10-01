#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""_verify_rubric.py — ab-003（case2）逐条 rubric 核验，直接检查 treatment 产物。

rubric 各维度：
  1. 负责人解析：'负责人是我，本周五之前出结果' 行 → 负责人='质检钱主任'（我→说话人）、期限='周五之前'
  2. 期限解析：'5月20日之前必须到港' 行 → 期限='5月20日之前'；'我们需要重新提交检测报告' 行 → 负责人/期限均='待定'
  3. 边界收录：14:18 '结论：……。散会' 行归入决议事项
  4. 条数：决议 3 / 待办 3 / 风险 3，summary.json 与 docx 一致且合计=9；待办为含 负责人/期限 列的表格
  5. 溯源与噪声：14:00 寒暄与 14:02/14:03/14:13 无关键词行未混入；每条目可在 case2.txt 逐字找到出处
附加：条目格式 <i>. 【hh:mm · 说话人】内容；与参照产物 oracle/out/case2/纪要.docx 文本级一致。
"""
import json
import os
import re
import sys

from docx import Document

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ASSET = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", ".."))
TREAT = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(TREAT, "纪要.docx")
SUMMARY = os.path.join(TREAT, "summary.json")
TRANSCRIPT = os.path.join(ASSET, "oracle", "inputs", "case2.txt")
REF_DOCX = os.path.join(ASSET, "oracle", "out", "case2", "纪要.docx")

SECTION_HEADS = ("一、决议事项", "二、待办事项", "三、风险与关注")
results = []


def check(name, ok, detail):
    results.append((name, bool(ok), detail))
    print("[%s] %s — %s" % ("PASS" if ok else "FAIL", name, detail))


doc = Document(DOCX)
paras = [p.text for p in doc.paragraphs]

# 按节标题切分段落
sections = {h: [] for h in SECTION_HEADS}
current = None
for p in paras:
    t = p.strip()
    if t in SECTION_HEADS:
        current = t
        continue
    if current is not None and t:
        sections[current].append(t)

# 待办表格
todo_table = None
for tb in doc.tables:
    if not tb.rows:
        continue
    header = [c.text.strip() for c in tb.rows[0].cells]
    if "负责人" in header and "期限" in header:
        todo_table = tb
        break
todo_rows = [(r.cells[0].text.strip(), r.cells[1].text.strip(), r.cells[2].text.strip())
             for r in todo_table.rows[1:]] if todo_table else []

# ---- rubric 1：负责人是我 → 质检钱主任；期限=周五之前 ----
row_owner = next((r for r in todo_rows if "负责人是我" in r[0]), None)
check("rubric1.owner_speaker_dealias", row_owner is not None and row_owner[1] == "质检钱主任",
      "事项='负责人是我，本周五之前出结果。' 行负责人=%r（期望 '质检钱主任'）" % (row_owner[1] if row_owner else None))
check("rubric1.deadline_friday", row_owner is not None and row_owner[2] == "周五之前",
      "同行的期限=%r（期望 '周五之前'，'本周' 前缀不进提取结果）" % (row_owner[2] if row_owner else None))

# ---- rubric 2：5月20日之前 / 待定+待定 ----
row_shipping = next((r for r in todo_rows if "5月20日之前" in r[0]), None)
check("rubric2.deadline_0520", row_shipping is not None and row_shipping[2] == "5月20日之前",
      "'需要空运，5月20日之前必须到港' 行期限=%r（期望 '5月20日之前'）" % (row_shipping[2] if row_shipping else None))
row_report = next((r for r in todo_rows if "重新提交检测报告" in r[0]), None)
check("rubric2.report_owner_pending", row_report is not None and row_report[1] == "待定",
      "'我们需要重新提交检测报告' 行负责人=%r（期望 '待定'）" % (row_report[1] if row_report else None))
check("rubric2.report_deadline_pending", row_report is not None and row_report[2] == "待定",
      "'我们需要重新提交检测报告' 行期限=%r（期望 '待定'）" % (row_report[2] if row_report else None))

# ---- rubric 3：14:18 收尾行含关键词 → 决议 ----
entry_1418 = next((t for t in sections["一、决议事项"] if t.startswith("【14:18")or "14:18" in t.split("】")[0]), None)
check("rubric3.boundary_1418_decision",
      entry_1418 is not None
      and "培训未完成的员工一律不许独立上岗。散会" in entry_1418,
      "决议节 14:18 条目=%r（收尾行含 '结论' 关键词，必须收录）" % entry_1418)

# ---- rubric 4：条数与 summary 一致 ----
n_dec = len([t for t in sections["一、决议事项"]])
n_risk = len([t for t in sections["三、风险与关注"]])
n_todo = len(todo_rows) if todo_table else 0
with open(SUMMARY, "r", encoding="utf-8-sig") as f:
    summary = json.load(f)
check("rubric4.counts_3_3_3", (n_dec, n_todo, n_risk) == (3, 3, 3),
      "docx 实际：决议 %d / 待办 %d / 风险 %d" % (n_dec, n_todo, n_risk))
check("rubric4.summary_matches_docx",
      summary.get("决议事项") == n_dec and summary.get("待办事项") == n_todo
      and summary.get("风险与关注") == n_risk
      and summary.get("合计") == n_dec + n_todo + n_risk == 9,
      "summary.json=%s（四键与 docx 一致，合计=9）" % json.dumps(summary, ensure_ascii=False))
check("rubric4.todo_table_owner_deadline_cols", todo_table is not None,
      "待办表格表头=%s，数据行 %d" % ([c.text.strip() for c in todo_table.rows[0].cells], n_todo))

# ---- rubric 5：噪声未混入 + 逐字溯源 ----
all_text = paras + [c.text for tb in doc.tables for row in tb.rows for c in row.cells]
all_text_j = "\n".join(all_text)
noise_lines = {
    "14:00 寒暄": "下午好，人都到了就开会",
    "14:02 无关键词": "2号线的主轴轴承坏了",
    "14:03 无关键词": "我已经安排夜班把1号线产能拉满",
    "14:13 无关键词": "德国客户要求7月15日前必须交付",
}
mixed = {k: (v in all_text_j) for k, v in noise_lines.items()}
check("rubric5.no_noise_lines", not any(mixed.values()), "噪声行混入情况=%s" % mixed)

with open(TRANSCRIPT, "r", encoding="utf-8-sig") as f:
    lines = f.read().splitlines()
bad = []
entries = []
ENTRY_PARSE = re.compile(r"^\s*\d+\.\s*【\s*([^·】]+?)\s*·\s*([^·】]+?)\s*】\s*(.*)$", re.S)
for t in sections["一、决议事项"] + sections["三、风险与关注"]:
    m = ENTRY_PARSE.match(t)
    if not m:
        bad.append("无法解析条目前缀: %s" % t[:30])
        continue
    entries.append((m.group(1).strip(), m.group(2).strip(), m.group(3).strip()))
for r in todo_rows:
    entries.append((None, None, r[0]))  # 待办表格按 runner 口径只溯源「事项」列
for time_s, speaker, content in entries:
    hit = next((ln for ln in lines if content in ln), None)
    if hit is None:
        bad.append("无出处: %s" % content[:30])
    elif time_s is not None and (("[%s]" % time_s) not in hit or speaker not in hit):
        bad.append("出处行时间/说话人不符: %s %s" % (time_s, speaker))
check("rubric5.verbatim_provenance", not bad,
      "共 %d 条目（决议/风险含时间+说话人校验，待办溯源「事项」列）全部在 case2.txt 逐字溯源%s"
      % (len(entries), "" if not bad else "；失败=%r" % bad))

# 条目格式 <i>. 【hh:mm · 说话人】内容
ENTRY_RE = re.compile(r"^\d+\.\s*【\d{1,2}:\d{2}\s*·\s*[^·】]+\s*】.+$", re.S)
fmt_bad = [t for t in sections["一、决议事项"] + sections["三、风险与关注"] if not ENTRY_RE.match(t)]
check("format.entry_pattern", not fmt_bad, "决议/风险条目格式 <i>. 【hh:mm · 说话人】内容 全部符合%s"
      % ("" if not fmt_bad else "；不符合=%r" % fmt_bad))

# 附加：与参照产物文本级一致
def docx_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return parts

check("extra.reference_text_identical",
      docx_text(DOCX) == docx_text(REF_DOCX),
      "treatment 纪要.docx 与参照 oracle/out/case2/纪要.docx 文本级一致")

n_fail = sum(1 for _, ok, _ in results if not ok)
print("\n%d/%d checks passed" % (len(results) - n_fail, len(results)))
sys.exit(0 if n_fail == 0 else 1)
