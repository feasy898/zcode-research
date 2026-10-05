#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""_verify_rubric_ab001.py — ab-001 rubric 专项校验（作用于顶层交付产物）

校验对象：<本目录>/纪要.docx 与 <本目录>/summary.json（顶层交付物）
依据：ab-001 rubric 5 维度 + SKILL.md 产物与验收表 + docx-and-summary.md 验收清单。
全部通过打印 PASS 并 exit 0；任一失败打印 FAIL 明细并 exit 1。
"""

import json
import os
import re
import sys

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))  # 资产根
DOCX = os.path.join(HERE, "纪要.docx")
SUMMARY = os.path.join(HERE, "summary.json")
TRANSCRIPT = os.path.join(ROOT, "oracle", "inputs", "case1.txt")
STAGING_DOCX = os.path.join(HERE, "_staging", "case1", "纪要.docx")
STAGING_SUMMARY = os.path.join(HERE, "_staging", "case1", "summary.json")

SECTION_HEADINGS = ("一、决议事项", "二、待办事项", "三、风险与关注")
ENTRY_RE = re.compile(r"^\s*\d+\.\s*【\s*([^·】]+?)\s*·\s*([^·】]+?)\s*】\s*(.*)$", re.S)
NOISE_SNIPPETS = [  # rubric 指定的 5 行噪声各自的独特片段
    ("09:00 寒暄行", "大家上午好，人到齐了我们就开始"),
    ("09:01 无关键词行", "上周服务压测结果出来了"),
    ("09:04 无关键词行", "大促推广费批了三百万"),
    ("09:12 口误更正行", "等等，我刚才说错了"),
    ("09:15 散会行", "好，今天先到这里，下周同一时间复盘"),
]

failures = []


def check(name, ok, detail):
    print("[%s] %s — %s" % ("PASS" if ok else "FAIL", name, detail))
    if not ok:
        failures.append((name, detail))


def docx_text(doc):
    """全部段落文本 + 全部表格单元格文本（顶层结构即验收对象）。"""
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return parts


with open(TRANSCRIPT, "r", encoding="utf-8-sig") as f:
    transcript_lines = f.read().splitlines()

# ---- 交付物存在性 ----
check("产物存在", os.path.isfile(DOCX) and os.path.isfile(SUMMARY),
      "%s / %s" % (os.path.basename(DOCX), os.path.basename(SUMMARY)))

doc = Document(DOCX)  # 打不开会抛异常 → 直接判 FAIL 结构维度
with open(SUMMARY, "r", encoding="utf-8-sig") as f:
    summary = json.load(f)

paras = [p.text for p in doc.paragraphs]

# ---- 维度1 结构：三节标题逐字精确且顺序正确 ----
pos = []
for h in SECTION_HEADINGS:
    hits = [i for i, t in enumerate(paras) if t.strip() == h]
    pos.append(hits[0] if hits else -1)
check("结构·三节标题齐全且顺序正确",
      pos == sorted(pos) and all(p >= 0 for p in pos),
      "段落位置 %s（要求 0<1<2 且全部存在，标题逐字=一、决议事项/二、待办事项/三、风险与关注）" % pos)

# ---- 维度2 条数：决议2/待办2/风险3，含'不能延期'行归风险；summary 四键一致、合计=7 ----
def section_entries(heading):
    out, started = [], False
    for t in paras:
        s = t.strip()
        if s == heading:
            started = True
            continue
        if started:
            if s in SECTION_HEADINGS:
                break
            if s:
                out.append(s)
    return out

dec = [t for t in section_entries("一、决议事项") if ENTRY_RE.match(t)]
risk = [t for t in section_entries("三、风险与关注") if ENTRY_RE.match(t)]
todo_table = next((t for t in doc.tables
                   if {"负责人", "期限"} <= {c.text.strip() for c in t.rows[0].cells}), None)
todo_rows = list(todo_table.rows[1:]) if todo_table else []

check("条数·决议=2", len(dec) == 2, "决议条目 %d 条" % len(dec))
check("条数·风险=3", len(risk) == 3, "风险条目 %d 条" % len(risk))
check("条数·'不能延期'行归风险",
      any("不能延期" in t for t in risk),
      "风险节含『48小时发货的承诺不能延期』行: %s" % any("48小时发货的承诺不能延期" in t for t in risk))
check("条数·summary四键",
      set(summary.keys()) == {"决议事项", "待办事项", "风险与关注", "合计"},
      "keys=%s" % sorted(summary.keys()))
check("条数·summary与docx一致",
      summary["决议事项"] == len(dec) and summary["风险与关注"] == len(risk)
      and summary["待办事项"] == len(todo_rows),
      "json(决议%s/待办%s/风险%s) vs docx(决议%d/待办%d/风险%d)"
      % (summary["决议事项"], summary["待办事项"], summary["风险与关注"],
         len(dec), len(todo_rows), len(risk)))
check("条数·合计=7=三者和",
      summary["合计"] == 7 and summary["合计"] == summary["决议事项"]
      + summary["待办事项"] + summary["风险与关注"],
      "合计=%s" % summary["合计"])
meta = next((t for t in paras if t.startswith("来源：")), "")
check("条数·元信息行与条数一致",
      "解析发言 7 条（决议 2 · 待办 2 · 风险 3）" in meta and "case1.txt" in meta,
      "元信息行: %s" % meta)

# ---- 维度3 待办表格：3 列、表头含 事项/负责人/期限、数据行 2 行 ----
header = [c.text.strip() for c in todo_table.rows[0].cells] if todo_table else []
check("表格·3列表头含事项/负责人/期限",
      todo_table is not None and len(todo_table.rows[0].cells) == 3
      and all(h in header for h in ("事项", "负责人", "期限")),
      "表头=%s" % header)
check("表格·数据行=2", len(todo_rows) == 2, "数据行 %d 行" % len(todo_rows))

# ---- 维度4 溯源：每条目【时间 · 说话人】前缀 + 原句逐字可溯源 ----
entries = []
for t in dec + risk:
    m = ENTRY_RE.match(t)
    if m:
        entries.append((m.group(1).strip(), m.group(2).strip(), m.group(3).strip(), t))
bad_prefix = [t for t in dec + risk if not ENTRY_RE.match(t)]
check("溯源·条目前缀格式", not bad_prefix,
      "决议+风险 %d 条均含【hh:mm · 说话人】前缀，不合格式 %d 条"
      % (len(dec) + len(risk), len(bad_prefix)))
bad_src = []
for when, speaker, content, _raw in entries:
    hit = next((ln for ln in transcript_lines if content in ln), None)
    if hit is None:
        bad_src.append("无出处:%s" % content[:20])
    elif ("[%s]" % when) not in hit or speaker not in hit:
        bad_src.append("时间/说话人不符:%s %s" % (when, speaker))
check("溯源·决议/风险逐字可溯", not bad_src and len(entries) == 5,
      "%d 条决议/风险条目全部溯源（%s）" % (len(entries), bad_src or "含时间+说话人"))
todo_bad = []
for r in todo_rows:
    item = r.cells[0].text.strip()
    hit = next((ln for ln in transcript_lines if item in ln), None)
    if hit is None:
        todo_bad.append("事项列无出处:%s" % item[:20])
check("溯源·待办事项列逐字可溯", not todo_bad,
      "2 行待办『事项』列均为原句逐字引用（%s）" % (todo_bad or "全部命中"))
# 待办负责人/期限：或原文可逐字找到，或为「待定」
cell_bad = []
for r in todo_rows:
    item = r.cells[0].text.strip()
    for label, cell in (("负责人", r.cells[1].text.strip()), ("期限", r.cells[2].text.strip())):
        if cell != "待定" and cell not in item:
            cell_bad.append("%s=%s 不在原句且非待定" % (label, cell))
check("溯源·负责人/期限列合规", not cell_bad,
      "两行负责人/期限 = %s（%s）"
      % ([[r.cells[1].text.strip(), r.cells[2].text.strip()] for r in todo_rows],
         cell_bad or "原文可寻或待定"))

# ---- 维度5 噪声抑制：5 行噪声片段不得出现在全文（段落+表格）----
all_text = "\n".join(docx_text(doc))
noise_hits = [(label, frag) for label, frag in NOISE_SNIPPETS if frag in all_text]
check("噪声·5 行噪声未混入", not noise_hits,
      "检查 09:00/09:01/09:04/09:12/09:15 片段 -> %s" % (noise_hits or "全文零命中"))

# ---- 附：顶层交付物与 _staging/case1（eval/runner.py 已验收对象）文本一致性 ----
if os.path.isfile(STAGING_DOCX):
    s_doc = Document(STAGING_DOCX)
    check("一致性·顶层与已验收staging文本一致",
          docx_text(doc) == docx_text(s_doc),
          "段落+表格文本逐字相同" )
    with open(STAGING_SUMMARY, "r", encoding="utf-8-sig") as f:
        s_sum = json.load(f)
    check("一致性·summary与staging一致", summary == s_sum, "四键数值相同")
else:
    print("[SKIP] staging 产物不存在，跳过一致性比对")

print("=" * 60)
if failures:
    print("RESULT: FAIL (%d 项)" % len(failures))
    sys.exit(1)
print("RESULT: PASS — ab-001 rubric 全维度通过（顶层交付产物）")
