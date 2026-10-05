# -*- coding: utf-8 -*-
"""618大促备战会 会议纪要生成器（三段式：决议 / 待办表格 / 风险）
条目内容逐字引用 case1.txt 原句；脚本内置溯源断言，任何一句在原转写稿中
找不到 `[时间] 说话人: 原句` 完整出处即报错退出，杜绝改写与编造。
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
SRC = HERE.parents[3] / "oracle" / "inputs" / "case1.txt"

# ---- 条目定义（time, speaker, verbatim）-----------------------------------
DECISIONS = [
    ("09:02", "王磊", "优惠券模块还有两个已知问题，我们内部的结论是必须在上线前修完。"),
    ("09:06", "主持人", "好，这个方案就议定下来，费用走大促专项。"),
]
# 待办表格行：(time, speaker, verbatim, 负责人, 期限)
TODOS = [
    ("09:05", "刘洋", "客服侧需要临时增加二十个坐席，我周五前把名单给到外包公司。", "刘洋", "周五前"),
    ("09:13", "王磊", "准确数字是下单接口7500，优惠券接口还有瓶颈，我需要再压一轮。", "王磊", "文中未明确"),
]
RISKS = [
    ("09:08", "陈静", "有个风险要提示，物流侧去年大促爆过仓，今年单量预计还要涨三成。"),
    ("09:09", "刘洋", "我补充一句，万一华东仓爆仓，能不能延期发货？"),
    ("09:10", "主持人", "不行，48小时发货的承诺不能延期，我们要提前把运力锁掉。"),
]

# ---- 溯源断言：每条必须能在原转写稿中逐字找到出处 ---------------------------
source_text = SRC.read_text(encoding="utf-8")
for t, s, body in DECISIONS + RISKS:
    line = f"[{t}] {s}: {body}"
    assert line in source_text, f"溯源失败（原句不在转写稿中）: {line}"
for t, s, body, _, _ in TODOS:
    line = f"[{t}] {s}: {body}"
    assert line in source_text, f"溯源失败（原句不在转写稿中）: {line}"

def prefix(t, s):
    return f"【{t} · {s}】"

# ---- 生成 docx -------------------------------------------------------------
doc = Document()
doc.core_properties.title = "618大促备战会 会议纪要"

title = doc.add_paragraph()
run = title.add_run("618大促备战会 会议纪要")
run.bold = True
run.font.size = Pt(16)

doc.add_heading("一、决议事项", level=1)
for t, s, body in DECISIONS:
    doc.add_paragraph(prefix(t, s) + body)

doc.add_heading("二、待办事项", level=1)
table = doc.add_table(rows=1 + len(TODOS), cols=3)
table.style = "Table Grid"
for i, h in enumerate(["事项", "负责人", "期限"]):
    cell = table.rows[0].cells[i]
    cell.text = h
    cell.paragraphs[0].runs[0].bold = True
for r, (t, s, body, owner, deadline) in enumerate(TODOS, start=1):
    table.rows[r].cells[0].text = prefix(t, s) + body
    table.rows[r].cells[1].text = owner
    table.rows[r].cells[2].text = deadline

doc.add_heading("三、风险与关注", level=1)
for t, s, body in RISKS:
    doc.add_paragraph(prefix(t, s) + body)

docx_path = HERE / "纪要.docx"
doc.save(docx_path)

# ---- 生成 summary.json（四键：三类条数 + 合计，与 docx 实际内容一致）--------
summary = {
    "决议事项": len(DECISIONS),
    "待办事项": len(TODOS),
    "风险与关注": len(RISKS),
    "合计": len(DECISIONS) + len(TODOS) + len(RISKS),
}
(HERE / "summary.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)

print("docx:", docx_path)
print("summary:", json.dumps(summary, ensure_ascii=False))
assert summary["合计"] == 7, "合计必须为 7"
print("BUILD OK")
