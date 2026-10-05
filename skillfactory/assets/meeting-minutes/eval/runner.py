#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""eval/runner.py — meeting-minutes 确定性评测器（机器可判，无随机/无网络）

用法：
    python eval/runner.py <被测输出目录> <参照输出目录>
    例：python eval/runner.py package/out oracle/out     # 评被测实现
        python eval/runner.py oracle/out oracle/out      # 自校验（应 exit 0）

- 用例清单读本脚本同目录的 golden.json（eval_inputs 的 case 列表）。
- 被测/参照目录下各含 <case>/ 子目录（纪要.docx + summary.json）。
- 每用例固定 5 项检查，全部通过 exit 0 并打印 JSON；任一失败 exit 1。
  1) docx_opens_with_three_sections            docx 可被 python-docx 打开且含三节标题
  2) todo_is_table_with_owner_deadline_columns 待办为表格且表头含 负责人/期限 列
  3) summary_counts_match_docx                 summary.json 条数与 docx 实际内容一致
  4) no_fabricated_entries                     每条目能在转写稿中逐字找到出处（含时间+说话人）
  5) entry_counts_within_tolerance_of_reference 与参照产物比，各类条数与合计差值 ≤2
"""

import json
import os
import re
import sys

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
GOLDEN_PATH = os.path.join(HERE, "golden.json")

DOCX_NAME = "纪要.docx"
SUMMARY_NAME = "summary.json"
SECTION_HEADINGS = ("一、决议事项", "二、待办事项", "三、风险与关注")
SUMMARY_KEYS = ("决议事项", "待办事项", "风险与关注", "合计")
TODO_TABLE_HEADER = ("事项", "负责人", "期限")
ENTRY_RE = re.compile(r"^\s*\d+\.\s*【\s*([^·】]+?)\s*·\s*([^·】]+?)\s*】\s*(.*)$", re.S)
COUNT_TOLERANCE = 2


def load_summary(path):
    """读取 summary.json；返回 (dict, err)。"""
    if not os.path.isfile(path):
        return None, "缺 %s" % SUMMARY_NAME
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except Exception as exc:  # noqa: BLE001 - 报告原始错误
        return None, "%s 解析失败: %s" % (SUMMARY_NAME, exc)
    if not isinstance(data, dict):
        return None, "%s 不是 JSON 对象" % SUMMARY_NAME
    return data, None


def section_entries(doc):
    """按节标题切分段落，返回 {节标题: [编号条目段落文本]}。"""
    sections = {h: [] for h in SECTION_HEADINGS}
    current = None
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in SECTION_HEADINGS:
            current = text
            continue
        if current is not None and text:
            sections[current].append(text)
    return sections


def parse_entry(text):
    """解析 '1. 【hh:mm · 说话人】内容'；返回 (time, speaker, content) 或 None。"""
    m = ENTRY_RE.match(text)
    if not m:
        return None
    return m.group(1).strip(), m.group(2).strip(), m.group(3).strip()


def find_todo_table(doc):
    """定位待办表格：表头行含 负责人 与 期限 的第一张表；返回 (table, header) 或 (None, None)。"""
    for table in doc.tables:
        if not table.rows:
            continue
        header = [c.text.strip() for c in table.rows[0].cells]
        if "负责人" in header and "期限" in header:
            return table, header
    return None, None


def check_case(case, cand_dir, ref_dir):
    """对单个 case 执行 5 项检查，返回 checks 列表（name/pass/detail）。"""
    checks = []
    cand_case = os.path.join(cand_dir, case)
    ref_case = os.path.join(ref_dir, case)
    docx_path = os.path.join(cand_case, DOCX_NAME)
    summary_path = os.path.join(cand_case, SUMMARY_NAME)

    def add(name, ok, detail):
        checks.append({"name": "%s/%s" % (case, name), "pass": bool(ok),
                       "detail": detail})
        return ok

    doc = None
    if not os.path.isdir(cand_case):
        pre = "被测用例目录不存在: %s" % cand_case
    elif not os.path.isfile(docx_path):
        pre = "缺 %s" % docx_path
    else:
        pre = None
    if pre is None:
        try:
            doc = Document(docx_path)
        except Exception as exc:  # noqa: BLE001 - 报告原始错误
            pre = "docx 无法被 python-docx 打开: %s" % exc
    if pre is not None:
        for name in ("docx_opens_with_three_sections",
                     "todo_is_table_with_owner_deadline_columns",
                     "summary_counts_match_docx",
                     "no_fabricated_entries",
                     "entry_counts_within_tolerance_of_reference"):
            add(name, False, pre)
        return checks

    # 1) 三节标题
    sections = section_entries(doc)
    missing = [h for h in SECTION_HEADINGS if not any(
        p.strip() == h for p in [para.text for para in doc.paragraphs])]
    add("docx_opens_with_three_sections", not missing,
        "三节标题齐全（一/二/三）" if not missing
        else "缺节标题: %s" % "、".join(missing))

    # 2) 待办表格与负责人/期限列
    todo_table, header = find_todo_table(doc)
    add("todo_is_table_with_owner_deadline_columns", todo_table is not None,
        "待办表格表头=%s（共 %d 数据行）" % (header, len(todo_table.rows) - 1)
        if todo_table is not None
        else "未找到表头含 负责人/期限 列的待办表格（文档共 %d 张表）" % len(doc.tables))

    # 3) summary.json 与 docx 实际内容一致
    summary, err = load_summary(summary_path)
    if err:
        add("summary_counts_match_docx", False, err)
    else:
        missing_keys = [k for k in SUMMARY_KEYS if k not in summary]
        if missing_keys:
            add("summary_counts_match_docx", False,
                "%s 缺键: %s" % (SUMMARY_NAME, "、".join(missing_keys)))
        else:
            decision_items = [t for t in sections["一、决议事项"] if parse_entry(t)]
            risk_items = [t for t in sections["三、风险与关注"] if parse_entry(t)]
            todo_rows = max(len(todo_table.rows) - 1, 0) if todo_table else 0
            actual = {"决议事项": len(decision_items),
                      "待办事项": todo_rows,
                      "风险与关注": len(risk_items)}
            mismatches = ["%s: json=%s docx=%s" % (k, summary[k], actual[k])
                          for k in ("决议事项", "待办事项", "风险与关注")
                          if summary[k] != actual[k]]
            sum_ok = summary["合计"] == sum(summary[k] for k in
                                           ("决议事项", "待办事项", "风险与关注"))
            if mismatches:
                add("summary_counts_match_docx", False,
                    "条数不一致 -> " + "; ".join(mismatches))
            elif not sum_ok:
                add("summary_counts_match_docx", False,
                    "合计=%s 不等于三类之和=%s"
                    % (summary["合计"],
                       sum(summary[k] for k in ("决议事项", "待办事项", "风险与关注"))))
            else:
                add("summary_counts_match_docx", True,
                    "json=%s 与 docx 实际一致（决议 %d 段、待办 %d 行、风险 %d 段）"
                    % (json.dumps(summary, ensure_ascii=False),
                       len(decision_items), todo_rows, len(risk_items)))

    # 4) 无编造：每条目可在转写稿中逐字找到出处
    #    转写稿按布局约定位于参照目录上一级的 inputs/ 下（<资产根>/oracle/inputs/<case>.txt）
    ref_root = os.path.dirname(os.path.abspath(ref_dir))
    transcript_path = os.path.join(ref_root, "inputs", case + ".txt")
    if not os.path.isfile(transcript_path):
        add("no_fabricated_entries", False, "无法定位转写稿: %s" % transcript_path)
    else:
        with open(transcript_path, "r", encoding="utf-8-sig") as f:
            lines = f.read().splitlines()
        entries = []
        for text in sections["一、决议事项"] + sections["三、风险与关注"]:
            parsed = parse_entry(text)
            if parsed:
                entries.append(parsed)
        if todo_table is not None:
            for row in todo_table.rows[1:]:
                content = row.cells[0].text.strip() if row.cells else ""
                if content:
                    entries.append((None, None, content))
        bad = []
        for time_s, speaker, content in entries:
            hit = next((ln for ln in lines if content in ln), None)
            if hit is None:
                bad.append("无出处: %s" % content[:30])
            elif time_s is not None and (("[%s]" % time_s) not in hit
                                         or speaker not in hit):
                bad.append("出处行时间/说话人不符: %s %s" % (time_s, speaker))
        add("no_fabricated_entries", not bad,
            "共 %d 条目全部可在 %s 中逐字溯源（含时间+说话人）" % (len(entries), transcript_path)
            if not bad else "%d/%d 条目溯源失败: %s"
            % (len(bad), len(entries), "; ".join(bad[:3])))

    # 5) 与参照产物条数差值 ≤ 2
    ref_summary, err = load_summary(os.path.join(ref_case, SUMMARY_NAME))
    if err:
        add("entry_counts_within_tolerance_of_reference", False,
            "参照产物不可用: %s" % err)
    else:
        diffs = {}
        over = []
        for k in SUMMARY_KEYS:
            if k in summary and k in ref_summary:
                diff = abs(summary[k] - ref_summary[k])
                diffs[k] = diff
                if diff > COUNT_TOLERANCE:
                    over.append("%s 差 %d（被测 %s vs 参照 %s）"
                                % (k, diff, summary[k], ref_summary[k]))
        if len(diffs) < len(SUMMARY_KEYS):
            add("entry_counts_within_tolerance_of_reference", False,
                "被测或参照 %s 缺键，无法完整比对" % SUMMARY_NAME)
        elif over:
            add("entry_counts_within_tolerance_of_reference", False,
                "超出容差±%d -> " % COUNT_TOLERANCE + "; ".join(over))
        else:
            add("entry_counts_within_tolerance_of_reference", True,
                "各类与合计差值均 ≤%d: %s" % (
                    COUNT_TOLERANCE,
                    json.dumps(diffs, ensure_ascii=False)))
    return checks


def main(argv):
    if len(argv) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    cand_dir, ref_dir = argv[1], argv[2]
    with open(GOLDEN_PATH, "r", encoding="utf-8") as f:
        golden = json.load(f)
    cases = [item["case"] for item in golden.get("eval_inputs", [])]
    if not cases:
        print("golden.json eval_inputs 为空", file=sys.stderr)
        return 2

    checks = []
    for case in cases:
        checks.extend(check_case(case, cand_dir, ref_dir))

    ok = all(c["pass"] for c in checks)
    print(json.dumps({"ok": ok, "checks": checks}, ensure_ascii=False, indent=2))
    return 0 if ok else 1


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001 - 非 tty 场景可能不支持
        pass
    sys.exit(main(sys.argv))
