#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""eval/runner.py — 月度汇报生成器（monthly-report-ppt）确定性评测 runner

用法：
  python runner.py <被测输出目录> <参照输出目录> [--inputs <源xlsx所在目录>]

例：
  python eval/runner.py package/out oracle/out     # 常规评测
  python eval/runner.py oracle/out oracle/out      # 自校验（应全过、exit 0）
  python eval/runner.py <空目录> oracle/out         # 红检（应报失败、exit 1）

产物单元定义：同时含 汇报.pptx 与 summary.json 的目录。
被测目录本身是单元 → 单单元评测；否则扫描其一级子目录（按名排序）。

检查项（对每个被测单元，name 前缀为单元相对名）：
  1. pptx_open_pages_ge4      pptx 可被 python-pptx 打开、页数≥4、summary.页数=实际页数
  2. charts_3png_and_embedded 3 张指定图表 PNG 齐全，且 pptx ≥2 页内嵌图片
  3. summary_matches_recalc   summary.json 与 runner 独立重算（从源 xlsx）一致，
                              数值相对容差 5%，集合键与字符串精确
  4. dirty_data_rules         去重恒等式、缺失日期不计入月度合计、缺失数量按 0 计
  5. pages_within_2_of_ref    与参照单元的 pptx 页数差 ≤2

输出：单个 JSON {"ok","tested","reference","checks":[{"name","pass","detail"}]}
      全部通过 exit 0；任一失败 exit 1（失败明细在 checks 中 pass=false 项）。

源 xlsx 定位（供检查 3/4 重算）：summary.json 的「输入文件」给出文件名，依次查
  --inputs 目录 → <被测根>/inputs → <被测根>/../inputs → <被测根>/../../inputs
  → <参照根>/inputs → <参照根>/../inputs → <参照根>/../../inputs
取第一个存在的同名文件。

确定性：本脚本无时间/随机源，同一输入两次运行输出一致。
重算部分为依据 spec.md 的独立实现，不 import oracle。
"""

import argparse
import json
import os
import sys
from datetime import date, datetime

PPTX_NAME = "汇报.pptx"
SUMMARY_NAME = "summary.json"
CHART_PNGS = ("月度趋势.png", "类别分布.png", "状态占比.png")
COLUMNS = ("日期", "事项", "类别", "数量", "状态")
COMPLETE_STATUSES = ("已完成", "完成")
DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y%m%d")
REL_TOL = 0.05          # 数值相对容差 5%
PAGE_MIN = 4            # pptx 页数下限
PAGE_DIFF_MAX = 2       # 与参照页数差上限
EMBEDDED_PIC_SLIDES_MIN = 2   # pptx 内嵌图片的页数下限（oracle 为 P2/P3 两页）


# ---------------------------------------------------------------- 工具

def close(a, b):
    """数值容差比较：相对 5%（附极小绝对下限，避免 0 除）。"""
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return False
    return abs(a - b) <= max(abs(b) * REL_TOL, 1e-9)


def read_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def find_units(root):
    """root 本身是单元 → [(".", root)]；否则扫描一级子目录（排序）。"""
    root = os.path.abspath(root)
    if not os.path.isdir(root):
        return []

    def is_unit(d):
        return (os.path.isfile(os.path.join(d, PPTX_NAME))
                and os.path.isfile(os.path.join(d, SUMMARY_NAME)))

    if is_unit(root):
        return [(".", root)]
    units = []
    for name in sorted(os.listdir(root)):
        d = os.path.join(root, name)
        if os.path.isdir(d) and is_unit(d):
            units.append((name, d))
    return units


def resolve_xlsx(name, inputs_dir, roots):
    """按 contract §5 的固定顺序定位源 xlsx。"""
    cands = []
    if inputs_dir:
        cands.append(os.path.join(inputs_dir, name))
    for base in roots:
        for rel in (("inputs",), ("..", "inputs"), ("..", "..", "inputs")):
            cands.append(os.path.join(base, *rel, name))
    for c in cands:
        if os.path.isfile(c):
            return os.path.abspath(c)
    return None


# ------------------------------------------------ 独立重算（不 import oracle）

def _txt(v):
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _pdate(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _txt(v)
    if not s:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _pqty(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = _txt(v).replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _numtxt(x):
    if abs(x - round(x)) < 1e-9:
        return str(int(round(x)))
    return ("%0.2f" % x).rstrip("0").rstrip(".")


def recalc(xlsx_path):
    """从源 xlsx 独立重算出与 summary.json 同口径的数字（spec §2/§3/§4）。"""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)

    header = None
    for row in it:
        cells = [_txt(c) for c in row]
        if any(cells):
            header = cells
            break
    if header is None:
        raise ValueError("第一个工作表没有表头行")
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise ValueError("表头缺少必需列: %s" % ",".join(missing))
    idx = {c: header.index(c) for c in COLUMNS}

    raw, blank = [], 0
    for row in it:
        vals = {c: (row[idx[c]] if idx[c] < len(row) else None) for c in COLUMNS}
        if all(_txt(vals[c]) == "" for c in COLUMNS):
            blank += 1
            continue
        raw.append(vals)
    wb.close()

    seen, kept, dup = set(), [], 0
    for r in raw:
        d, q = _pdate(r["日期"]), _pqty(r["数量"])
        key = (
            d.isoformat() if d is not None else _txt(r["日期"]),
            _txt(r["事项"]),
            _txt(r["类别"]),
            _numtxt(q) if q is not None else _txt(r["数量"]),
            _txt(r["状态"]),
        )
        if key in seen:
            dup += 1
            continue
        seen.add(key)
        kept.append(r)

    monthly, cats, sts = {}, {}, {}
    total = 0.0
    m_date = m_qty = comp = blocked = 0
    for r in kept:
        d, q = _pdate(r["日期"]), _pqty(r["数量"])
        if q is None:
            m_qty += 1
            q = 0.0
        total += q
        if d is None:
            m_date += 1
        else:
            m = d.strftime("%Y-%m")
            monthly[m] = monthly.get(m, 0.0) + q
        c = _txt(r["类别"]) or "未分类"
        cats[c] = cats.get(c, 0) + 1
        s = _txt(r["状态"]) or "未知"
        sts[s] = sts.get(s, 0) + 1
        if s in COMPLETE_STATUSES:
            comp += 1
        if "阻塞" in s:
            blocked += 1

    months = sorted(monthly)
    n = len(kept)
    top = sorted(cats, key=lambda c: (-cats[c], c))[0] if cats else "-"
    return {
        "原始行数": len(raw),
        "重复行数": dup,
        "去重后行数": n,
        "空行数": blank,
        "数量合计": total,
        "月份数": len(months),
        "统计区间": [months[0], months[-1]] if months else [],
        "月度数量合计": {m: monthly[m] for m in months},
        "类别数": len(cats),
        "类别条目数": cats,
        "状态数": len(sts),
        "状态条目数": sts,
        "已完成条数": comp,
        "完成率": round(comp * 100.0 / n, 1) if n else 0.0,
        "阻塞条数": blocked,
        "缺失日期行数": m_date,
        "缺失数量行数": m_qty,
        "条目最多类别": top,
    }


# ---------------------------------------------------------------- 检查

def inspect_pptx(path):
    """返回 (slides, pic_slides, pics, tables)；打不开抛异常。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    pic_slides = pics = tables = 0
    for s in prs.slides:
        p = sum(1 for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
        t = sum(1 for sh in s.shapes if getattr(sh, "has_table", False))
        pics += p
        tables += t
        if p:
            pic_slides += 1
    return len(prs.slides), pic_slides, pics, tables


def check_summary_numbers(summary, rec):
    """summary vs 独立重算。返回 (失败明细列表, 通过比较数)。"""
    bad, compared = [], 0

    # 顶层行数四项
    for k in ("原始行数", "重复行数", "去重后行数", "空行数"):
        if not close(summary.get(k), rec[k]):
            bad.append("%s: summary=%r 重算=%r" % (k, summary.get(k), rec[k]))
        else:
            compared += 1

    key = summary.get("关键汇总数字")
    if not isinstance(key, dict):
        return ["关键汇总数字 缺失或不是对象"], compared

    def num(k):
        nonlocal compared
        if not close(key.get(k), rec[k]):
            bad.append("关键汇总数字.%s: summary=%r 重算=%r"
                       % (k, key.get(k), rec[k]))
        else:
            compared += 1

    for k in ("数量合计", "月份数", "类别数", "状态数", "已完成条数",
              "完成率", "阻塞条数", "缺失日期行数", "缺失数量行数"):
        num(k)

    # 统计区间（列表精确）
    if list(key.get("统计区间") or []) != list(rec["统计区间"]):
        bad.append("关键汇总数字.统计区间: summary=%r 重算=%r"
                   % (key.get("统计区间"), rec["统计区间"]))
    else:
        compared += 1

    # 三个计数 dict：键集合精确、值带容差
    for k, rk in (("月度数量合计", "月度数量合计"),
                  ("类别条目数", "类别条目数"),
                  ("状态条目数", "状态条目数")):
        got = key.get(k) or {}
        want = rec[rk]
        if set(got) != set(want):
            bad.append("关键汇总数字.%s: 键不一致 summary=%s 重算=%s"
                       % (k, sorted(got), sorted(want)))
            continue
        for kk in want:
            if not close(got[kk], want[kk]):
                bad.append("关键汇总数字.%s[%s]: summary=%r 重算=%r"
                           % (k, kk, got[kk], want[kk]))
        compared += 1

    # 条目最多类别（字符串精确）
    if key.get("条目最多类别") != rec["条目最多类别"]:
        bad.append("关键汇总数字.条目最多类别: summary=%r 重算=%r"
                   % (key.get("条目最多类别"), rec["条目最多类别"]))
    else:
        compared += 1

    return bad, compared


def evaluate_unit(unit_name, unit_dir, ref_dir, inputs_dir, roots, checks):
    tag = unit_name if unit_name != "." else "root"

    def add(name, ok, detail):
        checks.append({"name": "%s/%s" % (tag, name), "pass": bool(ok),
                       "detail": detail})

    # 载入 summary
    try:
        summary = read_json(os.path.join(unit_dir, SUMMARY_NAME))
        if not isinstance(summary, dict):
            raise ValueError("summary.json 顶层不是对象")
    except Exception as e:
        for n in ("pptx_open_pages_ge4", "charts_3png_and_embedded",
                  "summary_matches_recalc", "dirty_data_rules",
                  "pages_within_2_of_ref"):
            add(n, False, "无法读取 summary.json: %s" % e)
        return

    # 1) pptx 打开与页数
    pptx_path = os.path.join(unit_dir, PPTX_NAME)
    slides = pic_slides = pics = tables = None
    try:
        slides, pic_slides, pics, tables = inspect_pptx(pptx_path)
    except Exception as e:
        add("pptx_open_pages_ge4", False, "python-pptx 无法打开 %s: %s"
            % (PPTX_NAME, e))
    else:
        ok = slides >= PAGE_MIN
        rep = summary.get("页数")
        consistent = (rep is None) or (rep == slides)
        detail = ("打开成功：实际页数=%d（要求≥%d），内嵌图片=%d（%d 页），表格=%d；"
                  "summary.页数=%s%s"
                  % (slides, PAGE_MIN, pics, pic_slides, tables, rep,
                     "" if consistent else "（与实际 %d 不一致）" % slides))
        add("pptx_open_pages_ge4", ok and consistent, detail)

    # 2) 图表齐全 + 已内嵌
    missing_png = [c for c in CHART_PNGS
                   if not os.path.isfile(os.path.join(unit_dir, c))]
    ok2 = (not missing_png) and slides is not None \
        and pic_slides >= EMBEDDED_PIC_SLIDES_MIN
    add("charts_3png_and_embedded", ok2,
        "图表 PNG 缺失=%s；pptx 含图片的页=%d（要求≥%d）"
        % (missing_png or "无", pic_slides if slides is not None else -1,
           EMBEDDED_PIC_SLIDES_MIN))

    # 3) summary vs 独立重算
    xlsx_name = summary.get("输入文件")
    xlsx = resolve_xlsx(xlsx_name, inputs_dir, roots) if xlsx_name else None
    if xlsx is None:
        add("summary_matches_recalc", False,
            "找不到源 xlsx（输入文件=%r，定位顺序见 contract §5；可用 --inputs 指定）"
            % xlsx_name)
    else:
        try:
            rec = recalc(xlsx)
        except Exception as e:
            add("summary_matches_recalc", False, "重算失败（%s）: %s" % (xlsx, e))
        else:
            bad, compared = check_summary_numbers(summary, rec)
            add("summary_matches_recalc", not bad,
                ("与独立重算一致：%d 项核对全部通过（源=%s，容差 5%%）"
                 % (compared, os.path.relpath(xlsx)))
                if not bad else
                ("与独立重算不一致 %d 项（源=%s，容差 5%%）：%s"
                 % (len(bad), os.path.relpath(xlsx), "；".join(bad[:8]))))

            # 4) 脏数据处理（依赖重算结果）
            sub = []
            raw_n, dup_n, ded_n = (summary.get("原始行数"),
                                   summary.get("重复行数"),
                                   summary.get("去重后行数"))
            if not (isinstance(raw_n, int) and isinstance(dup_n, int)
                    and isinstance(ded_n, int) and raw_n - dup_n == ded_n):
                sub.append("去重恒等式不成立：原始 %r − 重复 %r ≠ 去重后 %r"
                           % (raw_n, dup_n, ded_n))
            monthly_sum = sum((summary.get("关键汇总数字") or {})
                              .get("月度数量合计", {}).values())
            rec_monthly_sum = sum(rec["月度数量合计"].values())
            if not close(monthly_sum, rec_monthly_sum):
                sub.append("月度合计和 %r 与重算 %r 不一致"
                           % (monthly_sum, rec_monthly_sum))
            elif rec["缺失日期行数"] > 0 and not monthly_sum <= \
                    rec["数量合计"] + 1e-9:
                sub.append("缺失日期 %d 行却月度合计和 %r > 数量合计 %r"
                           % (rec["缺失日期行数"], monthly_sum,
                              rec["数量合计"]))
            if not close(summary.get("关键汇总数字", {}).get("缺失数量行数"),
                         rec["缺失数量行数"]):
                sub.append("缺失数量行数 summary=%r 重算=%d（缺失须按 0 计入合计）"
                           % (summary.get("关键汇总数字", {})
                              .get("缺失数量行数"), rec["缺失数量行数"]))
            add("dirty_data_rules", not sub,
                ("脏数据处理正确：去重 %d−%d=%d；缺失日期 %d 行未计入月度"
                 "（月度合计和=%s，数量合计=%s）；缺失数量 %d 行按 0 计"
                 % (raw_n, dup_n, ded_n, rec["缺失日期行数"], monthly_sum,
                    rec["数量合计"], rec["缺失数量行数"]))
                if not sub else "；".join(sub))

    # 5) 与参照页数差
    if ref_dir is None:
        add("pages_within_2_of_ref", False,
            "参照目录中未找到与 %r 对应的参照产物单元" % unit_name)
    else:
        try:
            r_slides = inspect_pptx(os.path.join(ref_dir, PPTX_NAME))[0]
        except Exception as e:
            add("pages_within_2_of_ref", False,
                "参照 pptx 无法打开（%s）: %s" % (ref_dir, e))
        else:
            diff = abs(slides - r_slides) if slides is not None else None
            ok5 = diff is not None and diff <= PAGE_DIFF_MAX
            add("pages_within_2_of_ref", ok5,
                "被测页数=%s，参照页数=%d，差=%s（允许≤%d）"
                % (slides, r_slides, diff, PAGE_DIFF_MAX))


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="monthly-report-ppt 确定性评测 runner")
    ap.add_argument("tested", help="被测输出目录")
    ap.add_argument("reference", help="参照输出目录")
    ap.add_argument("--inputs", default=None, help="源 xlsx 所在目录（可选）")
    args = ap.parse_args(argv)

    checks = []

    tested_units = find_units(args.tested)
    ref_units = find_units(args.reference)
    ref_by_rel = dict(ref_units)

    if not tested_units:
        checks.append({"name": "units_found", "pass": False,
                       "detail": "被测目录 %s 中未发现任何产物单元（单元=同时含 "
                                 "%s 与 %s 的目录）" % (args.tested, PPTX_NAME,
                                                       SUMMARY_NAME)})
    else:
        single_ref = ref_units[0][1] if len(ref_units) == 1 else None
        for rel, d in tested_units:
            try:
                s = read_json(os.path.join(d, SUMMARY_NAME))
            except Exception:
                s = {}
            ref_dir = ref_by_rel.get(rel)
            if ref_dir is None and single_ref is not None \
                    and single_ref != d:
                # 参照为单单元、被测为子目录时，仅在输入文件同名时才配对
                rs = {}
                try:
                    rs = read_json(os.path.join(single_ref, SUMMARY_NAME))
                except Exception:
                    pass
                if rs.get("输入文件") == s.get("输入文件"):
                    ref_dir = single_ref
            roots = [os.path.abspath(args.tested), os.path.abspath(args.reference),
                     os.path.dirname(d)]
            if ref_dir:
                roots.append(os.path.dirname(ref_dir))
            evaluate_unit(rel, d, ref_dir, args.inputs, roots, checks)

    checks.append({"name": "reference_dir_has_units", "pass": bool(ref_units),
                   "detail": "参照目录 %s 含 %d 个产物单元"
                             % (args.reference, len(ref_units))})

    ok = all(c["pass"] for c in checks)
    print(json.dumps({"ok": ok, "tested": args.tested,
                      "reference": args.reference, "checks": checks},
                     ensure_ascii=False, indent=2))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
