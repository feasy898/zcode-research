#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""oracle.py — 月度汇报 PPT 生成参照实现（oracle，确定性脚本）

输入：台账.xlsx（第一个工作表，表头行含：日期/事项/类别/数量/状态 五列；
     多余列忽略。日期列接受 date/datetime 单元格或字符串
     YYYY-MM-DD / YYYY/M/D / YYYY.M.D / YYYY年M月D日）。

清洗规则（全部确定性，按顺序执行）：
  1. 五列全空的行视为空行：不计入原始行数，单独记 空行数。
  2. 完全重复行（五列规范化后一致；数值列按数值比较）只保留首次出现。
     原始行数 = 非空数据行总数；去重后行数 = 原始行数 - 重复行数。
  3. 缺失值处理：
     日期 缺失/无法解析 → 该行不进月度趋势，计入 缺失日期行数；
     类别 缺失 → 归入「未分类」；状态 缺失 → 归入「未知」；
     数量 缺失/非数字 → 按 0 参与求和，计入 缺失数量行数。
  4. 「已完成/完成」视为完成态；状态含「阻塞」计为阻塞事项。

图表（matplotlib Agg 后端，中文字体 Microsoft YaHei，figsize/dpi 固定）：
  图1 月度趋势.png  各月数量合计折线图
  图2 类别分布.png  各类别条目数条形图
  图3 状态占比.png  各状态条目占比饼图

PPT（python-pptx，16:9，共 4 页，全部内容由数据规则推导，不含当前时间）：
  P1 封面；P2 数据概览（关键数字 + 嵌月度趋势图）；
  P3 类别明细（表格 + 嵌类别分布图）；P4 结论与下月计划。

用法：python oracle.py --input <台账.xlsx> --outdir <输出目录>
产物：<outdir>/月度趋势.png、类别分布.png、状态占比.png、汇报.pptx、summary.json
"""

import argparse
import json
import os
import sys
import time
import zipfile
from datetime import date, datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False

COLUMNS = ("日期", "事项", "类别", "数量", "状态")
COMPLETE_STATUSES = ("已完成", "完成")
FONT = "Microsoft YaHei"
PALETTE = ["#4472C4", "#ED7D31", "#A5A5A5", "#FFC000", "#5B9BD5",
           "#70AD47", "#264478", "#9E480E"]

DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y%m%d")


# ---------------------------------------------------------------- 读取与清洗

def _text(v):
    """单元格 → 规范化文本（None/空白 → 空串）。"""
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def parse_date(v):
    """日期列 → date；缺失/无法解析返回 None。"""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    if not s:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_qty(v):
    """数量列 → float；缺失/非数字返回 None（bool 不算数字）。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = _text(v).replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def fmt_num(x):
    """数值展示：整值去小数点，其余保留至多 2 位。"""
    if x is None:
        return "0"
    if abs(x - round(x)) < 1e-9:
        return str(int(round(x)))
    return ("%0.2f" % x).rstrip("0").rstrip(".")


def read_ledger(path):
    """读第一个工作表。返回 (rows, blank_rows)；
    rows 为 dict 列表（键为五列名），空行已剔除并单独计数。
    表头缺列时抛 ValueError。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = None
    for row in it:
        cells = [_text(c) for c in row]
        if any(cells):
            header = cells
            break
    if header is None:
        raise ValueError("工作表为空：没有表头行")
    idx = {}
    for col in COLUMNS:
        if col not in header:
            raise ValueError("表头缺少必需列：%s（实际表头：%s）" % (col, header))
        idx[col] = header.index(col)

    rows, blank = [], 0
    for row in it:
        vals = {col: row[idx[col]] if idx[col] < len(row) else None
                for col in COLUMNS}
        if all(_text(vals[c]) == "" for c in COLUMNS):
            blank += 1
            continue
        rows.append(vals)
    wb.close()
    return rows, blank


def dedupe(rows):
    """完全重复行只保留首次出现。返回 (kept, dup_count)。"""
    kept, seen, dup = [], set(), 0
    for r in rows:
        qty = parse_qty(r["数量"])
        key = (
            _text(r["日期"]) if parse_date(r["日期"]) is None
            else parse_date(r["日期"]).isoformat(),
            _text(r["事项"]),
            _text(r["类别"]),
            fmt_num(qty) if qty is not None else _text(r["数量"]),
            _text(r["状态"]),
        )
        if key in seen:
            dup += 1
            continue
        seen.add(key)
        kept.append(r)
    return kept, dup


# ---------------------------------------------------------------- 汇总

def aggregate(deduped):
    """对去重后的行做全部汇总，返回 dict（确定性：所有序列已排序）。"""
    monthly_qty, monthly_cnt = {}, {}
    cat_cnt, cat_qty = {}, {}
    st_cnt = {}
    total_qty = 0.0
    missing_date = missing_qty = completed = blocked = 0

    for r in deduped:
        d = parse_date(r["日期"])
        qty = parse_qty(r["数量"])
        if qty is None:
            missing_qty += 1
            qty = 0.0
        total_qty += qty

        if d is None:
            missing_date += 1
        else:
            m = d.strftime("%Y-%m")
            monthly_qty[m] = monthly_qty.get(m, 0.0) + qty
            monthly_cnt[m] = monthly_cnt.get(m, 0) + 1

        cat = _text(r["类别"]) or "未分类"
        cat_cnt[cat] = cat_cnt.get(cat, 0) + 1
        cat_qty[cat] = cat_qty.get(cat, 0.0) + qty

        st = _text(r["状态"]) or "未知"
        st_cnt[st] = st_cnt.get(st, 0) + 1
        if st in COMPLETE_STATUSES:
            completed += 1
        if "阻塞" in st:
            blocked += 1

    months = sorted(monthly_qty)
    cats = sorted(cat_cnt, key=lambda c: (-cat_cnt[c], c))
    statuses = sorted(st_cnt, key=lambda s: (-st_cnt[s], s))
    n = len(deduped)
    top_cat = cats[0] if cats else "-"
    return {
        "dedup_rows": n,
        "missing_date": missing_date,
        "missing_qty": missing_qty,
        "total_qty": total_qty,
        "months": months,
        "monthly_qty": [monthly_qty[m] for m in months],
        "monthly_cnt": [monthly_cnt[m] for m in months],
        "cats": cats,
        "cat_cnt": [cat_cnt[c] for c in cats],
        "cat_qty": [cat_qty[c] for c in cats],
        "statuses": statuses,
        "status_cnt": [st_cnt[s] for s in statuses],
        "completed": completed,
        "blocked": blocked,
        "top_cat": top_cat,
        "top_cat_cnt": cat_cnt.get(top_cat, 0),
    }


# ---------------------------------------------------------------- 图表

def _style_axes(ax):
    ax.grid(True, axis="y", linestyle="--", alpha=0.4)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)


def _empty_note(ax, msg):
    ax.text(0.5, 0.5, msg, ha="center", va="center", fontsize=14, color="#888888")
    ax.set_xticks([])
    ax.set_yticks([])


def chart_trend(months, sums, path):
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=150)
    if months:
        ax.plot(months, sums, marker="o", color="#2F5597", linewidth=2)
        for x, y in zip(months, sums):
            ax.annotate(fmt_num(y), (x, y), textcoords="offset points",
                        xytext=(0, 8), ha="center", fontsize=9)
        ax.set_xlabel("月份")
        ax.set_ylabel("数量合计")
    else:
        _empty_note(ax, "无有效日期数据")
    ax.set_title("月度数量趋势")
    _style_axes(ax)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def chart_category(cats, counts, path):
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=150)
    if cats:
        bars = ax.bar(cats, counts, color="#4472C4", width=0.55)
        for b, c in zip(bars, counts):
            ax.annotate(str(c),
                        (b.get_x() + b.get_width() / 2, b.get_height()),
                        textcoords="offset points", xytext=(0, 3),
                        ha="center", fontsize=10)
        ax.set_ylabel("条目数")
        ax.set_ylim(0, max(counts) * 1.18)
    else:
        _empty_note(ax, "无类别数据")
    ax.set_title("类别分布（条目数）")
    _style_axes(ax)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def chart_status(statuses, counts, path):
    fig, ax = plt.subplots(figsize=(6.4, 4.8), dpi=150)
    if statuses:
        colors = [PALETTE[i % len(PALETTE)] for i in range(len(statuses))]
        ax.pie(counts, labels=statuses, autopct="%1.1f%%", startangle=90,
               colors=colors, textprops={"fontsize": 10})
        ax.axis("equal")
    else:
        _empty_note(ax, "无状态数据")
    ax.set_title("状态占比（条目数）")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


# ---------------------------------------------------------------- PPT

def _set_run(run, size, bold, color):
    from pptx.util import Pt
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = FONT
    f.color.rgb = color
    # 中文 Eastern-Asia 字体需要写 rPr 的 a:ea/a:cs
    from lxml import etree
    from pptx.oxml.ns import qn
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = etree.SubElement(rPr, qn(tag))
        e.set("typeface", FONT)


def _add_text(slide, x, y, w, h, lines):
    """lines: [(text, size, bold, color, align), ...]"""
    from pptx.util import Inches
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, size, bold, color, align) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = text
        _set_run(run, size, bold, color)
    return tb


def _accent_title(slide, text):
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.5),
                                 Inches(0.12), Inches(0.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x2F, 0x55, 0x97)
    bar.line.fill.background()
    _add_text(slide, 0.85, 0.38, 10.5, 0.8,
              [(text, 30, True, RGBColor(0x1F, 0x38, 0x64), PP_ALIGN.LEFT)])


def _footer(slide, text):
    from pptx.util import Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    _add_text(slide, 0.6, 7.05, 12.1, 0.35,
              [(text, 9, False, RGBColor(0x8C, 0x8C, 0x8C), PP_ALIGN.LEFT)])


def build_conclusions(agg, dedup_rows, dup_rows, blank_rows, raw_rows):
    concl = []
    n = agg["dedup_rows"]
    span = ("%s ~ %s" % (agg["months"][0], agg["months"][-1])
            if agg["months"] else "无有效月份")
    concl.append("记录规模：原始 %d 行，剔除重复 %d 行、空行 %d 行后有效 %d 条，"
                 "数量合计 %s。" % (raw_rows, dup_rows, blank_rows, n,
                                   fmt_num(agg["total_qty"])))
    concl.append("覆盖 %d 个月（%s）；条目最多的类别为「%s」，共 %d 条。"
                 % (len(agg["months"]), span, agg["top_cat"], agg["top_cat_cnt"]))
    if len(agg["months"]) >= 2:
        first, last = agg["monthly_qty"][0], agg["monthly_qty"][-1]
        if first == 0:
            concl.append("趋势：末月（%s）数量合计 %s，首月为 0，绝对增量 %s。"
                         % (agg["months"][-1], fmt_num(last), fmt_num(last - first)))
        else:
            pct = (last - first) / abs(first) * 100.0
            word = "上升" if pct > 0 else ("下降" if pct < 0 else "持平")
            concl.append("趋势：末月（%s）数量合计 %s，较首月（%s）%s %.1f%%。"
                         % (agg["months"][-1], fmt_num(last),
                            agg["months"][0], word, abs(pct)))
    else:
        concl.append("趋势：统计区间不足两个月，无法计算趋势。")
    rate = (agg["completed"] / n * 100.0) if n else 0.0
    concl.append("完成情况：已完成 %d 条，完成率 %.1f%%；未完成 %d 条%s。"
                 % (agg["completed"], rate, n - agg["completed"],
                    ("，其中阻塞 %d 条" % agg["blocked"]) if agg["blocked"] else ""))
    concl.append("数据质量：日期缺失 %d 行（未计入趋势）、数量缺失 %d 行（按 0 计）。"
                 % (agg["missing_date"], agg["missing_qty"]))
    return concl


def build_plans(agg, months):
    last_month = months[-1] if months else "-"
    try:
        y, m = int(last_month[:4]), int(last_month[5:7])
        nxt = "%04d-%02d" % (y + (m // 12), (m % 12) + 1)
    except (ValueError, IndexError):
        nxt = "下月"
    plans = [
        "「%s」为当前最大投入类别（%d 条），下月（%s）保持资源投入并设定数量目标。"
        % (agg["top_cat"], agg["top_cat_cnt"], nxt),
    ]
    unfinished = agg["dedup_rows"] - agg["completed"]
    if agg["blocked"]:
        plans.append("优先清理阻塞事项（%d 条），明确责任人与解除条件。" % agg["blocked"])
    if unfinished:
        plans.append("对未完成事项（%d 条）逐项跟进，滚动更新状态。" % unfinished)
    if agg["missing_date"] or agg["missing_qty"]:
        plans.append("补齐台账质量：补录缺失日期 %d 行、缺失数量 %d 行。"
                     % (agg["missing_date"], agg["missing_qty"]))
    plans.append("沿用本报表口径（去重、缺失统计、趋势与占比）按月复盘。")
    return plans


def build_pptx(agg, src_name, dup_rows, blank_rows, raw_rows, charts, out_path):
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    DARK = RGBColor(0x1F, 0x38, 0x64)
    BODY = RGBColor(0x26, 0x26, 0x26)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    n = agg["dedup_rows"]
    span = ("%s ~ %s" % (agg["months"][0], agg["months"][-1])
            if agg["months"] else "无有效月份")
    rate = (agg["completed"] / n * 100.0) if n else 0.0

    # P1 封面
    s = prs.slides.add_slide(blank)
    _add_text(s, 1.0, 2.3, 11.3, 1.4,
              [("月度工作汇报", 48, True, DARK, PP_ALIGN.CENTER)])
    _add_text(s, 1.0, 3.9, 11.3, 1.2,
              [("数据来源：%s ｜ 统计区间：%s" % (src_name, span),
                16, False, BODY, PP_ALIGN.CENTER),
               ("有效记录 %d 条 · 数量合计 %s · 完成率 %.1f%%"
                % (n, fmt_num(agg["total_qty"]), rate),
                16, False, BODY, PP_ALIGN.CENTER)])
    _footer(s, "本演示由 oracle.py 确定性生成（openpyxl + matplotlib + python-pptx）")

    # P2 数据概览（关键数字 + 嵌月度趋势图）
    s = prs.slides.add_slide(blank)
    _accent_title(s, "数据概览")
    bullets = [
        ("原始行数 %d（重复 %d 行已剔除、空行 %d 行已跳过）"
         % (raw_rows, dup_rows, blank_rows), 14, False, BODY),
        ("去重后有效记录 %d 条，覆盖 %d 个月（%s）"
         % (n, len(agg["months"]), span), 14, False, BODY),
        ("数量合计 %s；条目最多类别：「%s」（%d 条）"
         % (fmt_num(agg["total_qty"]), agg["top_cat"], agg["top_cat_cnt"]),
         14, False, BODY),
        ("已完成 %d 条，完成率 %.1f%%" % (agg["completed"], rate), 14, False, BODY),
        ("数据缺失：日期 %d 行（未计入趋势）、数量 %d 行（按 0 计）"
         % (agg["missing_date"], agg["missing_qty"]), 14, False, BODY),
        ("附图 3 张：月度趋势 / 类别分布 / 状态占比", 14, False, BODY),
    ]
    _add_text(s, 0.7, 1.7, 5.6, 4.6,
              [(t, sz, b, c, PP_ALIGN.LEFT) for (t, sz, b, c) in bullets])
    s.shapes.add_picture(charts["trend"], Inches(6.5), Inches(1.8),
                         width=Inches(6.4))
    _footer(s, "数据来源：%s ｜ 口径：去重后行级统计" % src_name)

    # P3 类别明细（表格 + 嵌类别分布图）
    s = prs.slides.add_slide(blank)
    _accent_title(s, "类别明细")
    rows = len(agg["cats"]) + 2
    gf = s.shapes.add_table(rows, 4, Inches(0.6), Inches(1.7),
                            Inches(6.4), Inches(0.5 + 0.45 * rows))
    table = gf.table
    for c, w in enumerate((2.2, 1.3, 1.6, 1.3)):
        table.columns[c].width = Inches(w)
    header_bg = RGBColor(0xDC, 0xE6, 0xF2)

    def fill_cell(cell, text, bold=False, bg=None):
        cell.fill.solid()
        if bg is not None:
            cell.fill.fore_color.rgb = bg
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = text
        _set_run(run, 12, bold, DARK if bold else BODY)

    for j, h in enumerate(("类别", "条目数", "数量合计", "占比")):
        fill_cell(table.cell(0, j), h, bold=True, bg=header_bg)
    for i, cat in enumerate(agg["cats"], 1):
        cnt = agg["cat_cnt"][i - 1]
        fill_cell(table.cell(i, 0), cat)
        fill_cell(table.cell(i, 1), str(cnt))
        fill_cell(table.cell(i, 2), fmt_num(agg["cat_qty"][i - 1]))
        fill_cell(table.cell(i, 3),
                  ("%.1f%%" % (cnt * 100.0 / n)) if n else "-")
    fill_cell(table.cell(rows - 1, 0), "合计", bold=True, bg=header_bg)
    fill_cell(table.cell(rows - 1, 1), str(n), bold=True, bg=header_bg)
    fill_cell(table.cell(rows - 1, 2), fmt_num(agg["total_qty"]), bold=True,
              bg=header_bg)
    fill_cell(table.cell(rows - 1, 3), "100%" if n else "-", bold=True,
              bg=header_bg)
    s.shapes.add_picture(charts["category"], Inches(7.3), Inches(1.9),
                         width=Inches(5.6))
    _footer(s, "注：类别缺失的行归入「未分类」")

    # P4 结论与下月计划
    s = prs.slides.add_slide(blank)
    _accent_title(s, "结论与下月计划")
    concl = build_conclusions(agg, n, dup_rows, blank_rows, raw_rows)
    plans = build_plans(agg, agg["months"])
    left = [("结论", 18, True, DARK)] + [
        ("· " + t, 13, False, BODY) for t in concl]
    right = [("下月计划", 18, True, DARK)] + [
        ("· " + t, 13, False, BODY) for t in plans]
    _add_text(s, 0.7, 1.6, 6.0, 5.2,
              [(t, sz, b, c, PP_ALIGN.LEFT) for (t, sz, b, c) in left])
    _add_text(s, 7.0, 1.6, 5.8, 5.2,
              [(t, sz, b, c, PP_ALIGN.LEFT) for (t, sz, b, c) in right])
    _footer(s, "统计区间：%s ｜ 有效记录 %d 条" % (span, n))

    tmp = out_path + ".raw.tmp"
    prs.save(tmp)
    try:
        _write_normalized_zip(tmp, out_path)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    return len(prs.slides._sldIdLst)


def _write_normalized_zip(src, dst, attempts=4, delay=0.25):
    """把 src 的 zip 条目重写为固定时间戳后写到 dst（保证 pptx 字节级确定性）。

    Windows 上杀软/索引器会在文件刚创建后短暂持锁，故对 dst 的写入做少量重试；
    重试只影响时机，不影响产物内容。"""
    with zipfile.ZipFile(src, "r") as zin:
        entries = [(i, zin.read(i.filename)) for i in zin.infolist()]
    last_err = None
    for attempt in range(attempts):
        try:
            with open(dst, "wb") as f, \
                    zipfile.ZipFile(f, "w", zipfile.ZIP_DEFLATED) as zout:
                for info, data in entries:
                    ni = zipfile.ZipInfo(info.filename,
                                         date_time=(1980, 1, 1, 0, 0, 0))
                    ni.compress_type = zipfile.ZIP_DEFLATED
                    ni.external_attr = info.external_attr
                    zout.writestr(ni, data)
            return
        except PermissionError as e:
            last_err = e
            time.sleep(delay)
    raise last_err


# ---------------------------------------------------------------- 主流程

def main(argv=None):
    ap = argparse.ArgumentParser(description="月度汇报 PPT 参照实现（oracle）")
    ap.add_argument("--input", required=True, help="台账.xlsx（列：日期/事项/类别/数量/状态）")
    ap.add_argument("--outdir", required=True, help="输出目录")
    args = ap.parse_args(argv)

    if not os.path.isfile(args.input):
        print("输入文件不存在: %s" % args.input, file=sys.stderr)
        return 2
    os.makedirs(args.outdir, exist_ok=True)

    try:
        rows, blank_rows = read_ledger(args.input)
    except ValueError as e:
        print("输入不合法: %s" % e, file=sys.stderr)
        return 2
    deduped, dup_rows = dedupe(rows)
    agg = aggregate(deduped)

    p_trend = os.path.join(args.outdir, "月度趋势.png")
    p_cat = os.path.join(args.outdir, "类别分布.png")
    p_status = os.path.join(args.outdir, "状态占比.png")
    chart_trend(agg["months"], agg["monthly_qty"], p_trend)
    chart_category(agg["cats"], agg["cat_cnt"], p_cat)
    chart_status(agg["statuses"], agg["status_cnt"], p_status)

    pptx_path = os.path.join(args.outdir, "汇报.pptx")
    pages = build_pptx(agg, os.path.basename(args.input), dup_rows,
                       blank_rows, len(rows),
                       {"trend": p_trend, "category": p_cat, "status": p_status},
                       pptx_path)

    summary = {
        "输入文件": os.path.basename(args.input),
        "原始行数": len(rows),
        "重复行数": dup_rows,
        "去重后行数": agg["dedup_rows"],
        "空行数": blank_rows,
        "图数": 3,
        "页数": pages,
        "图表": ["月度趋势.png", "类别分布.png", "状态占比.png"],
        "关键汇总数字": {
            "数量合计": agg["total_qty"],
            "月份数": len(agg["months"]),
            "统计区间": [agg["months"][0], agg["months"][-1]] if agg["months"] else [],
            "月度数量合计": dict(zip(agg["months"], agg["monthly_qty"])),
            "类别数": len(agg["cats"]),
            "类别条目数": dict(zip(agg["cats"], agg["cat_cnt"])),
            "状态数": len(agg["statuses"]),
            "状态条目数": dict(zip(agg["statuses"], agg["status_cnt"])),
            "已完成条数": agg["completed"],
            "完成率": round(agg["completed"] * 100.0 / agg["dedup_rows"], 1)
                      if agg["dedup_rows"] else 0.0,
            "阻塞条数": agg["blocked"],
            "缺失日期行数": agg["missing_date"],
            "缺失数量行数": agg["missing_qty"],
            "条目最多类别": agg["top_cat"],
        },
    }
    json_path = os.path.join(args.outdir, "summary.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print("输入: %s（原始 %d 行，重复 %d，空行 %d，去重后 %d 行）"
          % (args.input, len(rows), dup_rows, blank_rows, agg["dedup_rows"]))
    print("汇总: 数量合计 %s ｜ 覆盖 %d 个月 ｜ 完成率 %.1f%%"
          % (fmt_num(agg["total_qty"]), len(agg["months"]),
             summary["关键汇总数字"]["完成率"]))
    for p in (p_trend, p_cat, p_status, pptx_path, json_path):
        print("产物: %s（%d 字节）" % (p, os.path.getsize(p)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
