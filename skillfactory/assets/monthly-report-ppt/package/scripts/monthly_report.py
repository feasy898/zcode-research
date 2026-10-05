#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""monthly_report.py — 月度汇报生成器（Excel 台账进、图表+PPT 出）唯一入口脚本。

命令行契约（contract.md §2，冻结）：
    python monthly_report.py --input <台账.xlsx> --outdir <输出目录>

退出码：
    0  成功（stdout 打印人读摘要）
    2  输入文件不存在 / 第一个工作表无表头行 / 表头缺任一必需列
       （stderr 中文报错，不产生产物、不写半成品）

产物（contract.md §3，冻结，5 个文件全部直接写入 <outdir> 根，无任何子目录嵌套）：
    月度趋势.png / 类别分布.png / 状态占比.png   3 张中文统计图表
    汇报.pptx                                    4 页 16:9 汇报
    summary.json                                 机器可读关键汇总数字

行为语义遵循 spec.md：§2 输入契约（I1-I6）、§3 清洗（R1-R6）、§4 汇总（S1-S5）、
§5 图表（C1-C3）、§6 PPT（P1-P6）、§7 summary 键集合、§8 错误、§9 确定性（D1-D2）。

确定性实现要点（spec D1/D2）：
    - 脚本内无当前时间、无随机源；全部 PPT 文案由数据推导；
    - matplotlib 固定 Agg 后端，PNG metadata 固定（无时间戳块）；
    - pptx core properties 固定为常量；保存后把 zip 全部条目时间戳归一化为
      1980-01-01 再落盘；写最终文件遇 PermissionError 有限次重试。
    同输入连跑两次，5 个产物逐字节一致。
"""

import argparse
import io
import json
import math
import os
import sys
import time
import zipfile
from datetime import date, datetime

# ---- matplotlib：Agg 后端 + 中文字体，必须在 pyplot 之前（spec C1/C2）----
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

matplotlib.rcParams["font.family"] = "sans-serif"
matplotlib.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
matplotlib.rcParams["axes.unicode_minus"] = False   # 负号正常显示（C2）

FONT = "Microsoft YaHei"
COLUMNS = ("日期", "事项", "类别", "数量", "状态")
DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y%m%d")
COMPLETE_STATUSES = ("已完成", "完成")   # R6：全等匹配
BLOCK_SUBSTR = "阻塞"                    # R6：子串匹配
UNCATEGORIZED = "未分类"                 # R4
UNKNOWN_STATUS = "未知"                  # R4

CHART_MONTH = "月度趋势.png"
CHART_CATS = "类别分布.png"
CHART_STATUS = "状态占比.png"
CHARTS = (CHART_MONTH, CHART_CATS, CHART_STATUS)
PPTX_NAME = "汇报.pptx"
SUMMARY_NAME = "summary.json"

PNG_META = {"Software": None}            # 固定 metadata，PNG 不含时间/软件戳
COLOR_LINE = "#2E6FBA"
COLOR_BAR = "#4C90D0"
COLOR_GRAY = "#888888"

TEXT_DARK = RGBColor(0x33, 0x33, 0x33)
TEXT_GRAY = RGBColor(0x80, 0x80, 0x80)
TEXT_TITLE = RGBColor(0x1F, 0x3B, 0x57)
HEAD_FILL = "2F5597"
TOTAL_FILL = "D9E2F3"


class InputError(Exception):
    """输入不合法（无表头/缺必需列等），main 捕获后 exit 2（契约 §2）。"""


# ================================================================ 解析归一化

def _txt(v):
    """I6 文本规范化：None→空串；date/datetime→YYYY-MM-DD；其余 str().strip()。"""
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _pdate(v):
    """I4 日期解析：date/datetime 直取；字符串按 5 种格式依次尝试；失败→None。"""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _txt(v)
    if not s:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _pqty(v):
    """I5 数量解析：bool 不算数字；int/float 直取；字符串去半角逗号按 float。

    空/非法/NaN/Inf → None（缺失数量，R5 按 0 计入合计）。
    """
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        q = float(v)
    else:
        s = _txt(v).replace(",", "")
        if not s:
            return None
        try:
            q = float(s)
        except ValueError:
            return None
    if math.isnan(q) or math.isinf(q):
        return None
    return q


def _numtxt(x):
    """R2 去重键的数量归一化文本：整数去小数点；小数至多 2 位去尾零。"""
    if abs(x - round(x)) < 1e-9:
        return str(int(round(x)))
    return ("%0.2f" % x).rstrip("0").rstrip(".")


def _fmt_qty(x):
    """展示用数量文本（图表标注/PPT 文案共用，保证口径一致）。"""
    return _numtxt(x)


def _as_num(x):
    """JSON 数值：整数值转 int，其余保留 6 位小数（消除浮点尾差）。"""
    if abs(x - round(x)) < 1e-9:
        return int(round(x))
    return round(x, 6)


# ================================================================ 读取与清洗

def read_rows(xlsx_path):
    """I1-I3 + R1：读第一个工作表（data_only/read_only），表头=第一个非空行。

    返回 (raw_rows, blank)。五列全空的行不计原始行数，单独计空行（R1）。
    表头行本身不计入任何行数统计（I2）。
    """
    try:
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception as e:
        raise InputError("输入文件无法作为 xlsx 打开：%s" % e)
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)

        header = None
        for row in it:
            cells = [_txt(c) for c in row]
            if any(cells):
                header = cells
                break
        if header is None:
            raise InputError("工作表为空：没有表头行")
        missing = [c for c in COLUMNS if c not in header]
        if missing:
            raise InputError("表头缺少必需列：%s（实际表头：%s）"
                             % ("、".join(missing), header))
        idx = {c: header.index(c) for c in COLUMNS}

        raw, blank = [], 0
        for row in it:
            vals = {c: (row[idx[c]] if idx[c] < len(row) else None)
                    for c in COLUMNS}
            if all(_txt(vals[c]) == "" for c in COLUMNS):
                blank += 1                      # R1
                continue
            raw.append(vals)
        return raw, blank
    finally:
        wb.close()


def dedup(raw):
    """R2：完全重复行去重，键相同只保留首次出现，其后各计 1 次重复。

    键 =（日期可解析→ISO YYYY-MM-DD，否则原文；事项文本；类别文本；
    数量可解析→归一化数值文本，否则原文；状态文本）。
    恒等式：去重后行数 = 原始行数 − 重复行数。
    """
    seen, kept, dup = set(), [], 0
    for r in raw:
        d, q = _pdate(r["日期"]), _pqty(r["数量"])
        key = (
            d.isoformat() if d is not None else _txt(r["日期"]),
            _txt(r["事项"]),
            _txt(r["类别"]),
            _numtxt(q) if q is not None else _txt(r["数量"]),
            _txt(r["状态"]),
        )
        if key in seen:
            dup += 1
            continue
        seen.add(key)
        kept.append(r)
    return kept, dup


def compute_stats(kept, blank, raw, dup):
    """R3-R6 + S1-S5：对去重后的行做全部统计。返回聚合 dict。"""
    monthly, cat_cnt, cat_qty, st_cnt = {}, {}, {}, {}
    total = 0.0
    m_date = m_qty = comp = blocked = 0
    for r in kept:
        d, q = _pdate(r["日期"]), _pqty(r["数量"])
        if q is None:                       # R5：缺失数量按 0 参与求和
            m_qty += 1
            q = 0.0
        total += q
        if d is None:                       # R3：缺失日期不进月度聚合
            m_date += 1
        else:
            m = d.strftime("%Y-%m")
            monthly[m] = monthly.get(m, 0.0) + q
        c = _txt(r["类别"]) or UNCATEGORIZED            # R4
        cat_cnt[c] = cat_cnt.get(c, 0) + 1
        cat_qty[c] = cat_qty.get(c, 0.0) + q
        s = _txt(r["状态"]) or UNKNOWN_STATUS           # R4
        st_cnt[s] = st_cnt.get(s, 0) + 1
        if s in COMPLETE_STATUSES:                      # R6：全等
            comp += 1
        if BLOCK_SUBSTR in s:                           # R6：子串
            blocked += 1

    months = sorted(monthly)                                     # S2 月份升序
    cat_order = sorted(cat_cnt, key=lambda c: (-cat_cnt[c], c))  # S3
    st_order = sorted(st_cnt, key=lambda s: (-st_cnt[s], s))     # S3
    n = len(kept)
    top = cat_order[0] if cat_order else "-"                     # S5
    rate = round(comp * 100.0 / n, 1) if n else 0.0              # S4
    return {
        "raw": raw, "dup": dup, "n": n, "blank": blank,
        "total": total, "monthly": monthly, "months": months,
        "cat_cnt": cat_cnt, "cat_qty": cat_qty, "cat_order": cat_order,
        "st_cnt": st_cnt, "st_order": st_order,
        "comp": comp, "rate": rate, "blocked": blocked,
        "m_date": m_date, "m_qty": m_qty, "top": top,
    }


# ================================================================ 图表（C1-C3）

def _empty_hint(ax, text):
    """C3 空数据回退：居中灰色提示文字，不报错。"""
    ax.axis("off")
    ax.text(0.5, 0.5, text, ha="center", va="center",
            fontsize=16, color=COLOR_GRAY, transform=ax.transAxes)


def _save_fig(fig, path):
    fig.tight_layout()
    fig.savefig(path, dpi=150, metadata=PNG_META)
    plt.close(fig)


def _rotate_labels(ax):
    ax.tick_params(axis="x", rotation=30)
    for lbl in ax.get_xticklabels():
        lbl.set_ha("right")


def chart_monthly(agg, path):
    """C1 月度趋势：各月数量合计折线 + 数值标注；无有效月份画提示。"""
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=150)
    months = agg["months"]
    if not months:
        _empty_hint(ax, "无有效日期数据")
    else:
        ys = [agg["monthly"][m] for m in months]
        ax.plot(months, ys, marker="o", color=COLOR_LINE, linewidth=2)
        for x, y in zip(months, ys):
            ax.annotate(_fmt_qty(y), (x, y), textcoords="offset points",
                        xytext=(0, 8), ha="center", fontsize=10)
        ax.set_title("月度数量合计趋势", fontsize=13)
        ax.set_xlabel("月份")
        ax.set_ylabel("数量合计")
        ax.grid(True, axis="y", linestyle="--", alpha=0.4)
        ax.margins(y=0.18)
        if len(months) > 6:
            _rotate_labels(ax)
    _save_fig(fig, path)


def chart_categories(agg, path):
    """C1 类别分布：各类别条目数条形图（S3 排序）+ 数值标注。"""
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=150)
    order = agg["cat_order"]
    if not order:
        _empty_hint(ax, "无类别数据")
    else:
        counts = [agg["cat_cnt"][c] for c in order]
        xs = list(range(len(order)))
        ax.bar(xs, counts, color=COLOR_BAR)
        for x, y in zip(xs, counts):
            ax.annotate(str(y), (x, y), textcoords="offset points",
                        xytext=(0, 3), ha="center", fontsize=10)
        ax.set_xticks(xs)
        ax.set_xticklabels(order)
        if len(order) > 6:
            _rotate_labels(ax)
        ax.set_title("各类别条目数分布", fontsize=13)
        ax.set_ylabel("条目数")
        ax.grid(True, axis="y", linestyle="--", alpha=0.4)
        ax.margins(y=0.15)
    _save_fig(fig, path)


def chart_status(agg, path):
    """C1 状态占比：各状态条目占比饼图（S3 排序，显示百分比）。"""
    fig, ax = plt.subplots(figsize=(6.4, 4.8), dpi=150)
    order = agg["st_order"]
    if not order:
        _empty_hint(ax, "无状态数据")
    else:
        counts = [agg["st_cnt"][s] for s in order]
        ax.pie(counts, labels=order, autopct="%1.1f%%",
               startangle=90, counterclock=False)
        ax.set_title("状态条目占比", fontsize=13)
    _save_fig(fig, path)


# ================================================================ PPT（P1-P6）

SLIDE_W_EMU = 12192000   # 13.333 in（16:9）
SLIDE_H_EMU = 6858000    # 7.5 in


def _style_run(run, size=14, bold=False, color=None, name=FONT):
    """设置字体：latin + East-Asia 都指向 Microsoft YaHei（spec P1）。"""
    f = run.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    if color is not None:
        f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = rPr.makeelement(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {})
        latin = rPr.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
        if latin is not None:
            latin.addnext(ea)
        else:
            rPr.append(ea)
    ea.set("typeface", name)


def _add_text(slide, x, y, w, h, blocks):
    """blocks: [(text, {size,bold,color,align,space_after}), ...]，返回文本框。"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (txt, opt) in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(opt.get("space_after", 6))
        if opt.get("align"):
            p.alignment = opt["align"]
        run = p.add_run()
        run.text = txt
        _style_run(run, size=opt.get("size", 14), bold=opt.get("bold", False),
                   color=opt.get("color"))
    return box


def _add_footer(slide, text, page_no, page_total):
    """P6：每页灰色小字页脚（数据来源/口径说明/页码）。"""
    _add_text(slide, 0.5, 7.08, 12.3, 0.34, [
        ("%s ｜ 第 %d 页 / 共 %d 页" % (text, page_no, page_total),
         {"size": 9, "color": TEXT_GRAY}),
    ])


def _add_title(slide, text):
    _add_text(slide, 0.5, 0.32, 12.3, 0.8,
              [(text, {"size": 26, "bold": True, "color": TEXT_TITLE})])


def _range_text(agg):
    if not agg["months"]:
        return "无有效月份数据"
    return "%s ~ %s（%d 个月）" % (agg["months"][0], agg["months"][-1],
                                  len(agg["months"]))


def _trend_conclusion(agg):
    """P4 第 3 条：趋势（≥2 月：末月 vs 首月；首月 0：绝对增量；<2 月：明示）。"""
    months = agg["months"]
    if len(months) < 2:
        reason = "仅覆盖 %d 个月" % len(months) if months else "无有效月份数据"
        return "趋势：%s，无法计算趋势" % reason
    first_m, last_m = months[0], months[-1]
    f, l = agg["monthly"][first_m], agg["monthly"][last_m]
    if f == 0:
        return ("趋势：首月 %s 数量合计为 0，末月 %s 绝对增量 +%s"
                % (first_m, last_m, _fmt_qty(l)))
    pct = round((l - f) / f * 100.0, 1)
    if pct == 0:
        word = "持平"
    else:
        word = "上升" if pct > 0 else "下降"
    return ("趋势：末月 %s（%s）较首月 %s（%s）%s %s%%"
            % (last_m, _fmt_qty(l), first_m, _fmt_qty(f), word,
               _fmt_qty(abs(pct))))


def _next_month(agg):
    """P5：下月 = 数据最末月的次月（12 月 → 次年 1 月）。"""
    if not agg["months"]:
        return None
    y, m = agg["months"][-1].split("-")
    y, m = int(y), int(m)
    return "%04d-01" % (y + 1) if m == 12 else "%04d-%02d" % (y, m + 1)


def _conclusions(agg):
    """P4：结论固定 5 条（记录规模/覆盖与最多类别/趋势/完成情况/数据质量）。"""
    undone = agg["n"] - agg["comp"]
    rate_txt = _fmt_qty(agg["rate"])
    return [
        "记录规模：原始 %d 条，剔除重复 %d 条、空行 %d 条，有效 %d 条，"
        "数量合计 %s" % (agg["raw"], agg["dup"], agg["blank"],
                         agg["n"], _fmt_qty(agg["total"])),
        "覆盖 %s，条目最多类别为「%s」" % (_range_text(agg), agg["top"]),
        _trend_conclusion(agg),
        "完成情况：已完成 %d 条（完成率 %s%%），未完成 %d 条，其中阻塞 %d 条"
        % (agg["comp"], rate_txt, undone, agg["blocked"]),
        "数据质量：缺失日期 %d 行、缺失数量 %d 行（缺失数量按 0 计入合计，"
        "缺失日期不计入月度趋势）" % (agg["m_date"], agg["m_qty"]),
    ]


def _plans(agg):
    """P5：下月计划规则（最大类别保持投入→清阻塞→跟未完成→补缺失→固定末条）。"""
    undone = agg["n"] - agg["comp"]
    plans = []
    if agg["cat_order"]:
        plans.append("最大投入类别「%s」保持投入" % agg["top"])
    if agg["blocked"]:
        plans.append("优先清理 %d 条阻塞事项" % agg["blocked"])
    if undone:
        plans.append("%d 条未完成事项逐项跟进" % undone)
    if agg["m_date"] or agg["m_qty"]:
        plans.append("补录缺失数据：缺失日期 %d 行、缺失数量 %d 行"
                     % (agg["m_date"], agg["m_qty"]))
    plans.append("沿用本报表口径按月复盘")
    return plans


def _set_cell(cell, text, size=12, bold=False, color=None, fill=None,
              align=PP_ALIGN.CENTER):
    p = cell.text_frame.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run()
    run.text = text
    _style_run(run, size=size, bold=bold, color=color)
    p.alignment = align
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def build_pptx(agg, src_name, out_path):
    """P1-P6：16:9 恰 4 页（封面/数据概览/类别明细/结论与下月计划），空白版式自绘。

    返回实际页数（summary.页数 与之严格一致，eval 检查项 1）。
    """
    from pptx import Presentation

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank = prs.slide_layouts[6]

    footer = ("数据来源：%s ｜ 口径：剔除空行与完全重复行；缺失数量按 0 计，"
              "缺失日期不计入月度趋势；完成=状态恰为「已完成/完成」" % src_name)
    rate_txt = _fmt_qty(agg["rate"])

    # ---- P1 封面 ------------------------------------------------------
    s = prs.slides.add_slide(blank)
    _add_text(s, 1.0, 2.0, 11.3, 1.2,
              [("月度工作汇报", {"size": 44, "bold": True,
                                 "align": PP_ALIGN.CENTER,
                                 "color": TEXT_TITLE})])
    _add_text(s, 1.0, 3.5, 11.3, 2.4, [
        ("数据来源：%s" % src_name,
         {"size": 18, "align": PP_ALIGN.CENTER, "space_after": 10}),
        ("统计区间：%s" % _range_text(agg),
         {"size": 18, "align": PP_ALIGN.CENTER, "space_after": 10}),
        ("有效记录 %d 条 ｜ 数量合计 %s ｜ 完成率 %s%%"
         % (agg["n"], _fmt_qty(agg["total"]), rate_txt),
         {"size": 18, "align": PP_ALIGN.CENTER}),
    ])
    _add_footer(s, footer, 1, 4)

    # ---- P2 数据概览（内嵌月度趋势图）----------------------------------
    s = prs.slides.add_slide(blank)
    _add_title(s, "数据概览")
    bullets = [
        "行数：原始 %d 条，剔除重复 %d 条、空行 %d 条，去重后有效 %d 条"
        % (agg["raw"], agg["dup"], agg["blank"], agg["n"]),
        "覆盖月份：%s" % _range_text(agg),
        "数量合计：%s" % _fmt_qty(agg["total"]),
        "条目最多类别：%s" % agg["top"],
        "完成率：%s%%（已完成 %d / %d）" % (rate_txt, agg["comp"], agg["n"]),
        "缺失说明：缺失日期 %d 行（不计入月度趋势）；缺失数量 %d 行（按 0 计入合计）"
        % (agg["m_date"], agg["m_qty"]),
    ]
    _add_text(s, 0.55, 1.3, 6.6, 5.6,
              [("• " + b, {"size": 15, "space_after": 12}) for b in bullets])
    s.shapes.add_picture(os.path.join(os.path.dirname(out_path), CHART_MONTH),
                         Inches(7.35), Inches(1.6), width=Inches(5.4))
    _add_footer(s, footer, 2, 4)

    # ---- P3 类别明细（表格 + 内嵌类别分布图）---------------------------
    s = prs.slides.add_slide(blank)
    _add_title(s, "类别明细")
    order = agg["cat_order"]
    n_rows = len(order) + 2                    # 表头 + 数据行 + 合计行
    gf = s.shapes.add_table(n_rows, 4, Inches(0.55), Inches(1.4),
                            Inches(6.9), Inches(0.45) * n_rows)
    tbl = gf.table
    for j, w in enumerate((2.6, 1.3, 1.5, 1.5)):
        tbl.columns[j].width = Inches(w)
    for j, h in enumerate(("类别", "条目数", "数量合计", "占比")):
        _set_cell(tbl.cell(0, j), h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                  fill=HEAD_FILL)
    for i, c in enumerate(order, start=1):
        share = round(agg["cat_cnt"][c] * 100.0 / agg["n"], 1) if agg["n"] else 0.0
        _set_cell(tbl.cell(i, 0), c, align=PP_ALIGN.LEFT)
        _set_cell(tbl.cell(i, 1), str(agg["cat_cnt"][c]))
        _set_cell(tbl.cell(i, 2), _fmt_qty(agg["cat_qty"][c]))
        _set_cell(tbl.cell(i, 3), "%s%%" % _fmt_qty(share))
    last = n_rows - 1
    _set_cell(tbl.cell(last, 0), "合计", bold=True, fill=TOTAL_FILL)
    _set_cell(tbl.cell(last, 1), str(agg["n"]), bold=True, fill=TOTAL_FILL)
    _set_cell(tbl.cell(last, 2), _fmt_qty(agg["total"]), bold=True,
              fill=TOTAL_FILL)
    _set_cell(tbl.cell(last, 3), "100%" if agg["n"] else "-", bold=True,
              fill=TOTAL_FILL)
    s.shapes.add_picture(os.path.join(os.path.dirname(out_path), CHART_CATS),
                         Inches(7.7), Inches(1.7), width=Inches(5.1))
    _add_footer(s, footer, 3, 4)

    # ---- P4 结论与下月计划（左右两栏）----------------------------------
    s = prs.slides.add_slide(blank)
    _add_title(s, "结论与下月计划")
    _add_text(s, 0.55, 1.3, 6.15, 5.6,
              [("结论", {"size": 17, "bold": True, "space_after": 10})]
              + [("• " + t, {"size": 13, "space_after": 9})
                 for t in _conclusions(agg)])
    nxt = _next_month(agg)
    plan_title = ("下月计划（%s）" % nxt) if nxt else "下月计划"
    _add_text(s, 6.95, 1.3, 5.85, 5.6,
              [(plan_title, {"size": 17, "bold": True, "space_after": 10})]
              + [("• " + t, {"size": 13, "space_after": 9})
                 for t in _plans(agg)])
    _add_footer(s, footer, 4, 4)

    # ---- 确定性收尾（spec D2）------------------------------------------
    cp = prs.core_properties
    cp.author = "monthly-report-ppt"
    cp.last_modified_by = "monthly-report-ppt"
    cp.title = "月度工作汇报"
    cp.subject = ""
    cp.keywords = ""
    cp.comments = ""
    cp.revision = 1
    cp.created = datetime(1980, 1, 1, 0, 0, 0)
    cp.modified = datetime(1980, 1, 1, 0, 0, 0)

    buf = io.BytesIO()
    prs.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin:
        entries = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 0
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)
    data = payload.getvalue()

    attempts = 5
    for i in range(attempts):                  # PermissionError 有限重试（D2）
        try:
            with open(out_path, "wb") as fobj:
                fobj.write(data)
            break
        except PermissionError:
            if i == attempts - 1:
                raise
            time.sleep(0.3)
    return len(prs.slides)


# ================================================================ summary.json

def build_summary(agg, src_name, pages):
    """spec §7：顶层 9 键 + 关键汇总数字 14 键，键集合恰好等于契约 §3。"""
    months = agg["months"]
    return {
        "输入文件": src_name,
        "原始行数": agg["raw"],
        "重复行数": agg["dup"],
        "去重后行数": agg["n"],
        "空行数": agg["blank"],
        "图数": len(CHARTS),
        "页数": pages,
        "图表": list(CHARTS),
        "关键汇总数字": {
            "数量合计": _as_num(agg["total"]),
            "月份数": len(months),
            "统计区间": [months[0], months[-1]] if months else [],
            "月度数量合计": {m: _as_num(agg["monthly"][m]) for m in months},
            "类别数": len(agg["cat_cnt"]),
            "类别条目数": {c: agg["cat_cnt"][c] for c in agg["cat_order"]},
            "状态数": len(agg["st_cnt"]),
            "状态条目数": {s: agg["st_cnt"][s] for s in agg["st_order"]},
            "已完成条数": agg["comp"],
            "完成率": agg["rate"],
            "阻塞条数": agg["blocked"],
            "缺失日期行数": agg["m_date"],
            "缺失数量行数": agg["m_qty"],
            "条目最多类别": agg["top"],
        },
    }


# ================================================================ 主流程

def run(input_path, outdir):
    """成功路径：读取→清洗→统计→图表→PPT→summary；返回 0。"""
    if not os.path.isfile(input_path):
        raise InputError("输入文件不存在: %s" % input_path)
    raw, blank = read_rows(input_path)     # E2 两类错误在此抛出（未建目录前）
    kept, dup = dedup(raw)
    agg = compute_stats(kept, blank, len(raw), dup)
    src_name = os.path.basename(input_path)

    os.makedirs(outdir, exist_ok=True)     # 不存在递归创建，已存在复用（不清空）

    # 3 张图表：产物在 outdir 根（契约 §3 冻结文件名，不建任何子目录）
    chart_monthly(agg, os.path.join(outdir, CHART_MONTH))
    chart_categories(agg, os.path.join(outdir, CHART_CATS))
    chart_status(agg, os.path.join(outdir, CHART_STATUS))

    pages = build_pptx(agg, src_name, os.path.join(outdir, PPTX_NAME))

    summary = build_summary(agg, src_name, pages)
    with open(os.path.join(outdir, SUMMARY_NAME), "w", encoding="utf-8",
              newline="\n") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
        f.write("\n")

    k = summary["关键汇总数字"]
    print("[monthly-report-ppt] 处理完成：%s" % src_name)
    print("  行数：原始 %d → 剔除重复 %d → 有效 %d（空行 %d）"
          % (summary["原始行数"], summary["重复行数"],
             summary["去重后行数"], summary["空行数"]))
    print("  汇总：数量合计 %s；完成率 %s%%；条目最多类别 %s；统计区间 %s"
          % (_fmt_qty(float(k["数量合计"])), _fmt_qty(agg["rate"]),
             k["条目最多类别"], _range_text(agg)))
    for name in CHARTS + (PPTX_NAME, SUMMARY_NAME):
        print("  产物：%s" % os.path.join(outdir, name))
    return 0


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")   # 控制台编码兜底，不影响产物
        except Exception:
            pass
    ap = argparse.ArgumentParser(
        description="月度汇报生成器：Excel 台账 → 3 张图表 + 4 页 PPT + summary.json")
    ap.add_argument("--input", required=True, help="台账 xlsx 路径（第一个工作表）")
    ap.add_argument("--outdir", required=True, help="输出目录（不存在则递归创建）")
    args = ap.parse_args(argv)
    try:
        return run(args.input, args.outdir)
    except InputError as e:
        print(str(e), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
